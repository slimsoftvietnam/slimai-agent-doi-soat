"""Dashboard cho kết quả đối soát: trang Tóm tắt trong Excel và tệp HTML tự chứa.

Nhận dữ liệu đã đối soát từ agent.py; không đọc ảnh, không tính lại trạng thái.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
import json
from pathlib import Path
import re

STATUS_ORDER = ("Đã trả đủ", "Trả thiếu", "Chưa trả", "Cần kiểm tra")
# Bảng màu trạng thái cố định: tốt / nghiêm trọng vừa / nghiêm trọng / cảnh báo.
STATUS_HEX = {"Đã trả đủ": "0CA30C", "Trả thiếu": "EC835A", "Chưa trả": "D03B3B", "Cần kiểm tra": "FAB219"}
STATUS_TINT = {"Đã trả đủ": "E3F4E3", "Trả thiếu": "FCEAE2", "Chưa trả": "F8E1E1", "Cần kiểm tra": "FEF3DB"}
PROOF_KINDS = ("Đã ghép", "Gợi ý", "Chưa ghép", "Không đọc được", "Trùng")
TEMPLATE = Path(__file__).with_name("dashboard.html")


def _hour(text) -> int | None:
    match = re.search(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)", str(text or ""))
    return int(match.group(1)) if match else None


def _day(text) -> date | None:
    for d, m, y in re.findall(r"(?<!\d)(\d{1,2})/(\d{1,2})/(\d{4})", str(text or "")):
        try:
            return date(int(y), int(m), int(d))
        except ValueError:
            continue
    return None


def build_report(orders: list[dict], proofs: list[dict], *, large_order: float, source_images: str,
                 source_orders: str, invoices: list[dict] | None = None) -> dict:
    """orders: rows already carrying status/action/received/missing. proofs: rows carrying kind."""
    days = Counter(d for d in (_day(p.get("transferred_at")) for p in proofs) if d)
    report_day = days.most_common(1)[0][0] if days else date.today()

    totals = {
        "due": sum(o["due"] or 0 for o in orders),
        "received": sum(o["received"] for o in orders),
        "outstanding": sum(o["missing"] or 0 for o in orders),
        "surplus": sum(max(o["received"] - (o["due"] or 0), 0) for o in orders if o["due"] is not None),
        "surplus_count": sum(1 for o in orders if o["due"] is not None and o["received"] > o["due"] + 1),
    }
    by_status = {s: {"count": 0, "due": 0.0, "received": 0.0, "missing": 0.0} for s in STATUS_ORDER}
    for o in orders:
        bucket = by_status[o["status"]]
        bucket["count"] += 1
        bucket["due"] += o["due"] or 0
        bucket["received"] += o["received"]
        bucket["missing"] += o["missing"] or 0

    kinds = Counter(p["kind"] for p in proofs)
    usable = len(proofs) - kinds["Trùng"]
    hourly = defaultdict(lambda: {"amount": 0.0, "count": 0})
    for p in proofs:
        if p["kind"] == "Đã ghép" and p.get("amount") and _hour(p.get("transferred_at")) is not None:
            slot = hourly[_hour(p.get("transferred_at"))]
            slot["amount"] += p["amount"]
            slot["count"] += 1

    no_hour = [p for p in proofs if p["kind"] == "Đã ghép" and p.get("amount") and _hour(p.get("transferred_at")) is None]

    channels = defaultdict(lambda: {"due": 0.0, "received": 0.0, "count": 0})
    for o in orders:
        if o.get("channel"):
            c = channels[str(o["channel"])]
            c["due"] += o["due"] or 0
            c["received"] += min(o["received"], o["due"] or 0)
            c["count"] += 1

    return {
        "title": "Đối soát chuyển khoản",
        "date": report_day.strftime("%d/%m/%Y"),
        "generated": datetime.now().strftime("%H:%M %d/%m/%Y"),
        "sources": {"images": source_images, "orders": source_orders},
        "large_order": large_order,
        "totals": totals,
        "status": [{"name": s, **by_status[s]} for s in STATUS_ORDER],
        "ready": by_status["Đã trả đủ"]["count"],
        "todo": len(orders) - by_status["Đã trả đủ"]["count"],
        "proofs": {
            "total": len(proofs),
            "kinds": {k: kinds[k] for k in PROOF_KINDS},
            "auto_rate": kinds["Đã ghép"] / usable if usable else 0,
            "duplicate_amount": sum(p.get("amount") or 0 for p in proofs if p["kind"] == "Trùng"),
            "unassigned_amount": sum(p.get("amount") or 0 for p in proofs if p["kind"] in ("Gợi ý", "Chưa ghép")),
        },
        "hourly": [{"hour": h, **hourly[h]} for h in sorted(hourly)],
        "hourly_missing": {"count": len(no_hour), "amount": sum(p["amount"] for p in no_hour)},
        "channels": [{"name": k, **v} for k, v in sorted(channels.items(), key=lambda kv: -kv[1]["due"])],
        "large": [o for o in orders if o["status"] == "Đã trả đủ" and (o["due"] or 0) >= large_order],
        "orders": orders,
        "images": [{k: p.get(k) for k in ("source_file", "amount", "transferred_at", "matched_order", "result", "kind")} for p in proofs],
        "invoices": invoices,
    }


def read_invoice_reconciliation(path: Path) -> list[dict]:
    """Read the 'Đối chiếu' sheet written by the doi-chieu command."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Đối chiếu"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    items = []
    for row in rows[1:]:
        status = str(row[-1] or "")
        kind = ("Khớp" if status == "Khớp" else "Trùng" if status.startswith("Khóa trùng")
                else "Một bên" if status.startswith("Chỉ có") else "Lệch")
        # doi-chieu ghi "bảng trái" = tệp hóa đơn đã đọc, "bảng phải" = bảng kê phần mềm; đổi sang lời dễ hiểu.
        readable = (status.replace("Khóa trùng, không ghép duy nhất", "Tệp hóa đơn bị tải trùng")
                    .replace("Thiếu dữ liệu: ", "Thiếu "))
        # Hạ chữ đầu của tên cột ("Tổng tiền" → "tổng tiền") nhưng giữ nguyên viết tắt như "MST".
        readable = re.sub(r"Sai khác: (\w)(\w)", lambda m: "Lệch " + (m.group(1).lower() if m.group(2).islower() else m.group(1)) + m.group(2), readable)
        readable = re.sub(r"Chỉ có ở bảng trái \([^)]*\)", "Có tệp hóa đơn, phần mềm chưa ghi nhận", readable)
        readable = re.sub(r"Chỉ có ở bảng phải \([^)]*\)", "Có trong phần mềm, thiếu tệp hóa đơn", readable)
        items.append({"number": row[0], "status": readable, "kind": kind})
    return items


# ---------------- Excel ----------------

# Bảng màu dùng chung với dashboard HTML (xem DASHBOARD_STYLE.md).
PAGE, CARD, SOFT, INK, INK2, MUTED, LINE = "EEF1F5", "FFFFFF", "F4F6F9", "0D0F12", "5B6170", "9097A3", "ECEEF2"
LIME, LIME_INK, ORANGE, PURPLE, UP, DOWN, WARN = "C8F26D", "1F2A06", "F4772E", "9B87F5", "16A34A", "E5484D", "C48A00"
CARD_COLS = (("B", "D"), ("F", "H"), ("J", "L"), ("N", "P"))


def key_points(report: dict) -> tuple[list[dict], list[dict]]:
    """Bốn số chính và các ý chính, cùng nội dung với dashboard HTML."""
    t, st = report["totals"], {s["name"]: s for s in report["status"]}
    orders, n = report["orders"], len(report["orders"])
    rate = t["received"] / t["due"] if t["due"] else 0
    remind = [o for o in orders if o["status"] in ("Trả thiếu", "Chưa trả")]
    remind_sum = sum(o["missing"] or 0 for o in remind)
    big_sum = sum(o["due"] or 0 for o in report["large"])
    kpis = [
        {"label": "Tổng giá trị đơn", "icon": "◫", "value": t["due"], "fmt": "money", "tone": UP, "foot": f"▲ {n} đơn hôm nay"},
        {"label": "Còn phải thu", "icon": "△", "value": t["outstanding"], "fmt": "money", "tone": DOWN, "foot": f"▼ {report['todo']} đơn chưa xong"},
        {"label": "Đơn sẵn sàng giao", "icon": "⛟", "value": f"{report['ready']}/{n}", "fmt": "text", "tone": UP,
         "foot": f"▲ {_vnd(st['Đã trả đủ']['due'])} đủ tiền"},
        {"label": "Khách cần nhắc", "icon": "☏", "value": len(remind), "fmt": "count", "tone": DOWN if remind else UP,
         "foot": f"▼ {_vnd(remind_sum)} cần đòi"},
    ]
    insights = [
        {"tone": UP, "title": f"Thu {rate:.0%}", "text": f"{_vnd(t['received'])} trên {_vnd(t['due'])}; {report['ready']} đơn đủ tiền có thể giao."},
        {"tone": DOWN, "title": f"{report['todo']} đơn chưa xong",
         "text": f"Còn {_vnd(t['outstanding'])}: {st['Trả thiếu']['count']} thiếu, {st['Chưa trả']['count']} chưa thấy ảnh, "
                 f"{st['Cần kiểm tra']['count']} cần mở ảnh."},
    ]
    if report["large"]:
        insights.append({"tone": WARN, "title": f"{len(report['large'])} đơn lớn chờ sao kê",
                         "text": f"Trị giá {_vnd(big_sum)}; xác nhận tiền về rồi mới giao."})
    if report.get("invoices"):
        kinds = Counter(i["kind"] for i in report["invoices"])
        issues = len(report["invoices"]) - kinds["Khớp"]
        insights.append({"tone": DOWN if issues else UP, "title": f"{issues} hóa đơn cần xem" if issues else "Hóa đơn khớp hết",
                         "text": f"{kinds['Lệch']} lệch, {kinds['Trùng']} trùng, {kinds['Một bên']} một bên; {kinds['Khớp']} khớp."})
    return kpis, insights


def _vnd(value) -> str:
    return f"{value or 0:,.0f}".replace(",", ".") + "đ"


def excel_dashboard(ws, report: dict) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.drawing.fill import ColorChoice, PatternFillProperties
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = ws.parent
    ws.sheet_view.showGridLines = False
    widths = {"A": 3, "E": 2.5, "I": 2.5, "M": 2.5, "Q": 3}
    for i in range(1, 18):
        col = chr(64 + i)
        ws.column_dimensions[col].width = widths.get(col, 12.5)

    def fill(hex_):
        return PatternFill("solid", fgColor="FF" + hex_)

    def put(cell, value, size=10, bold=False, color=INK, align="left", fmt=None, wrap=False, bg=None):
        c = ws[cell]
        c.value = value
        c.font = Font(name="Arial", size=size, bold=bold, color="FF" + color)
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fmt:
            c.number_format = fmt
        if bg:
            c.fill = fill(bg)
        return c

    def block(first, last, top, bottom, color=CARD):
        for r in range(top, bottom + 1):
            for c in range(ord(first), ord(last) + 1):
                ws[f"{chr(c)}{r}"].fill = fill(color)

    def money_fmt(cell):
        cell.number_format = '#,##0"đ"'

    t = report["totals"]
    rate = t["received"] / t["due"] if t["due"] else 0
    kpis, insights = key_points(report)
    todo = [o for s in ("Trả thiếu", "Chưa trả", "Cần kiểm tra") for o in report["orders"] if o["status"] == s]
    last_row = 70 + len(todo) + len(report["status"])
    block("A", "Q", 1, last_row, PAGE)

    # Thanh trên cùng: logo · menu viên thuốc · báo cáo
    ws.row_dimensions[2].height = 30
    put("B2", "▤  Đối soát", 15, True)
    for col, label, active in (("F", "Tổng quan", True), ("G", "Việc cần làm", False), ("H", "Đơn hàng", False), ("J", "Hóa đơn", False)):
        put(f"{col}2", label, 9.5, active, LIME_INK if active else INK2, "center", bg=LIME if active else CARD)
    ws.merge_cells("N2:P2")
    put("N2", f"Báo cáo cuối ngày · {report['date']}", 9.5, True, INK, "center", bg=CARD)

    # Lời chào
    ws.row_dimensions[4].height = 28
    put("B4", "Đối soát cuối ngày 👋", 18, True)
    put("B5", f"{len(report['orders'])} đơn · {report['proofs']['total']} ảnh chuyển khoản · {report['date']} · "
              "Hôm nay thu được bao nhiêu, còn ai chưa trả, đơn nào giao được.", 9.5, False, INK2)

    # Bốn thẻ số chính
    for (first, last), k in zip(CARD_COLS, kpis):
        block(first, last, 7, 11)
        put(f"{first}7", k["label"], 10, False, INK2)
        put(f"{last}7", k["icon"], 11, False, INK, "right")
        ws.merge_cells(f"{first}8:{last}9")
        v = put(f"{first}8", k["value"], 22, True)
        if k["fmt"] == "money":
            money_fmt(v)
        put(f"{first}10", k["foot"], 9, True, k["tone"])
        put(f"{last}11", "Xem chi tiết ↗", 8.5, False, MUTED, "right")
    for r, h in ((7, 22), (8, 20), (9, 20), (10, 18), (11, 16)):
        ws.row_dimensions[r].height = h

    # Ghi chú
    ws.merge_cells("B13:P13")
    put("B13", "⚠  Số tiền lấy từ ảnh khách gửi, chưa phải xác nhận tiền đã vào tài khoản. Xem sao kê trước khi giao đơn lớn. "
               "Agent không tự gửi tin nhắn.", 9, False, INK2, bg=CARD)
    ws.row_dimensions[13].height = 22

    # Dữ liệu biểu đồ để ở trang ẩn
    data = wb.create_sheet("_bieu_do")
    data.sheet_state = "hidden"
    hours = report["hourly"]
    lo = min([8] + [h["hour"] for h in hours]) if hours else 8
    hi = max([18] + [h["hour"] for h in hours]) if hours else 18
    by_hour = {h["hour"]: h for h in hours}
    peak = max(hours, key=lambda h: h["amount"])["hour"] if hours else None
    data.append(["Giờ", "Tiền về", "Cộng dồn"])
    acc = 0
    for h in range(lo, hi + 1):
        acc += by_hour.get(h, {}).get("amount", 0)
        data.append([f"{h}h", by_hour.get(h, {}).get("amount", 0), acc])
    data.append([])
    status_row = data.max_row + 1
    data.append(["Trạng thái", "Giá trị đơn"])
    for s in report["status"]:
        data.append([s["name"], s["due"]])

    # Thẻ "Tổng quan thu tiền": số lớn + biểu đồ cột gạch chéo, giờ cao điểm màu cam
    block("B", "L", 15, 33)
    put("B15", "Tổng quan thu tiền", 11, True)
    put("L15", "···", 11, True, INK2, "right")
    ws.merge_cells("B16:D17")
    v = put("B16", t["received"], 22, True)
    money_fmt(v)
    put("E16", f"{rate:.0%} ↗", 9, True, LIME_INK, "center", bg=LIME)
    put("H16", "Phải thu", 8.5, False, MUTED)
    put("H17", _vnd(t["due"]), 10.5, True)
    put("K16", "Còn phải thu", 8.5, False, MUTED)
    put("K17", _vnd(t["outstanding"]), 10.5, True)
    put("B18", "▨ Tiền về theo giờ    ● Giờ tiền về nhiều nhất", 8.5, False, INK2)
    if hours:
        bars = BarChart()
        bars.type, bars.style, bars.legend, bars.gapWidth = "col", 2, None, 45
        bars.y_axis.numFmt = '#,##0'
        bars.y_axis.majorGridlines.spPr = None
        bars.add_data(Reference(data, min_col=2, min_row=1, max_row=hi - lo + 2), titles_from_data=True)
        bars.set_categories(Reference(data, min_col=1, min_row=2, max_row=hi - lo + 2))
        series = bars.series[0]
        series.graphicalProperties.pattFill = PatternFillProperties(
            prst="wdUpDiag", fgClr=ColorChoice(srgbClr="D9DCE2"), bgClr=ColorChoice(srgbClr="F1F2F5"))
        series.graphicalProperties.line.solidFill = "E1E3E8"
        if peak is not None:
            point = DataPoint(idx=peak - lo)
            point.graphicalProperties.solidFill = ORANGE
            point.graphicalProperties.line.solidFill = ORANGE
            series.dPt.append(point)
        bars.height, bars.width = 7.2, 21.5
        ws.add_chart(bars, "B19")

    # Thẻ "Ý chính hôm nay": mỗi ý một ô nền nhạt
    block("N", "P", 15, 33)
    put("N15", "Ý chính hôm nay", 11, True)
    put("P15", "···", 11, True, INK2, "right")
    put("N16", "Những điều cần biết trước khi chốt ngày", 8.5, False, MUTED)
    row = 18
    for item in insights[:4]:
        block("N", "P", row, row + 2, SOFT)
        ws.merge_cells(f"N{row}:P{row}")
        put(f"N{row}", "● " + item["title"], 11, True, item["tone"], bg=SOFT)
        ws.merge_cells(f"N{row + 1}:P{row + 2}")
        put(f"N{row + 1}", item["text"], 8.5, False, INK2, wrap=True, bg=SOFT).alignment = Alignment(wrap_text=True, vertical="top")
        row += 4

    # Thẻ "Giao dịch trong ngày" (đường tím cộng dồn) + "Giá trị đơn theo trạng thái"
    block("B", "H", 35, 50)
    put("B35", "Giao dịch trong ngày", 11, True)
    put("H35", "···", 11, True, INK2, "right")
    tx = sum(h["count"] for h in hours)
    ws.row_dimensions[36].height = 30
    put("B36", f"{tx} giao dịch", 20, True)
    if peak is not None:
        put("E36", f"Cao điểm {peak}h", 9, True, LIME_INK, "center", bg=LIME)
    if hours:
        line = LineChart()
        line.style, line.legend = 2, None
        line.y_axis.numFmt = '#,##0'
        line.add_data(Reference(data, min_col=3, min_row=1, max_row=hi - lo + 2), titles_from_data=True)
        line.set_categories(Reference(data, min_col=1, min_row=2, max_row=hi - lo + 2))
        s = line.series[0]
        s.smooth = True
        s.graphicalProperties.line.solidFill = PURPLE
        s.graphicalProperties.line.width = 32000
        line.height, line.width = 5.6, 13.6
        ws.add_chart(line, "B38")

    block("J", "P", 35, 50)
    put("J35", "Giá trị đơn theo trạng thái", 11, True)
    put("P35", "···", 11, True, INK2, "right")
    for i, st in enumerate(report["status"], start=37):
        put(f"J{i}", "● " + st["name"], 10, True, STATUS_HEX[st["name"]])
        put(f"N{i}", st["count"], 9, False, MUTED, "right").number_format = '0" đơn"'
        put(f"P{i}", st["due"], 10, True, INK, "right").number_format = '#,##0"đ"'
    status_chart = BarChart()
    status_chart.type, status_chart.style, status_chart.legend, status_chart.gapWidth = "bar", 2, None, 55
    status_chart.y_axis.numFmt = '#,##0'
    status_chart.add_data(Reference(data, min_col=2, min_row=status_row, max_row=status_row + 4), titles_from_data=True)
    status_chart.set_categories(Reference(data, min_col=1, min_row=status_row + 1, max_row=status_row + 4))
    for idx, st in enumerate(report["status"]):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = STATUS_HEX[st["name"]]
        point.graphicalProperties.line.noFill = True
        status_chart.series[0].dPt.append(point)
    status_chart.height, status_chart.width = 4.2, 13.6
    ws.add_chart(status_chart, "J42")

    # Thẻ "Việc cần làm hôm nay"
    r = 52
    block("B", "P", r, r + 3 + max(len(todo), 1))
    put(f"B{r}", "Việc cần làm hôm nay", 11, True)
    put(f"P{r}", "···", 11, True, INK2, "right")
    put(f"B{r + 1}", "Nhắc tiền còn thiếu trước, rồi đơn chưa thấy ảnh, rồi đơn cần mở ảnh", 8.5, False, MUTED)
    for col, h in zip("BCDGHN", ("Trạng thái", "Khách hàng", "", "Mã đơn · SĐT", "Việc cần làm", "Còn thiếu")):
        put(f"{col}{r + 2}", h, 8.5, False, MUTED, "right" if h == "Còn thiếu" else "left")
    row = r + 3
    for i, o in enumerate(todo):
        shade = SOFT if i % 2 == 0 else CARD
        block("B", "P", row, row, shade)
        put(f"B{row}", o["status"], 9, True, STATUS_HEX[o["status"]], bg=STATUS_TINT[o["status"]])
        put(f"C{row}", o["customer"], 10, True, bg=shade)
        put(f"G{row}", f"{o['code']} · {o['phone'] or ''}", 9, False, MUTED, bg=shade)
        put(f"H{row}", o["action"], 9, False, INK2, bg=shade)
        v = put(f"P{row}", o["missing"] or 0, 10, True, INK, "right", bg=shade)
        money_fmt(v)
        ws.row_dimensions[row].height = 22
        row += 1
    if not todo:
        put(f"B{row}", "Không có việc tồn. Mọi đơn đã đủ theo ảnh.", 10, False, MUTED)

    p = report["proofs"]
    put(f"B{row + 2}", (f"Ảnh: {p['total']} · ghép tự động {p['kinds']['Đã ghép']} ({p['auto_rate']:.0%}) · "
                        f"trùng {p['kinds']['Trùng']} · tạo lúc {report['generated']}"), 8.5, False, MUTED)

    ws.sheet_properties.tabColor = LIME
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:Q{row + 3}"


# ---------------- HTML ----------------

def html_dashboard(path: Path, report: dict) -> Path:
    payload = json.dumps(report, ensure_ascii=False, default=str).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__REPORT__*/null", payload)
    path.write_text(html, encoding="utf-8")
    return path
