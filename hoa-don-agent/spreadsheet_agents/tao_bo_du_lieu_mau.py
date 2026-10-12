"""Tạo bộ dữ liệu mẫu cho agent đối soát chuyển khoản và hóa đơn điện tử.

Mọi tên ngân hàng, ví, doanh nghiệp, MST, số tài khoản và khách hàng đều là giả lập.
Chạy lại (cần font Arial/Times của macOS): python tao_bo_du_lieu_mau.py <thư-mục-đích>
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil
import sys

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont

FONT_DIR = Path("/System/Library/Fonts/Supplemental")
ARIAL, ARIAL_B = FONT_DIR / "Arial.ttf", FONT_DIR / "Arial Bold.ttf"
TIMES, TIMES_B = FONT_DIR / "Times New Roman.ttf", FONT_DIR / "Times New Roman Bold.ttf"
SHOP = "CONG TY TNHH HOA MOC DEMO"
SHOP_ACCOUNT = "0071 0001 23456"
NOTE = "Ảnh mẫu giả lập · Khóa học AI Agent SlimAI"

# Đơn hàng xuất từ phần mềm: mã, giờ, khách, SĐT, kênh, tiền hàng, giảm giá, khách cần trả, số HĐ
ORDERS = [
    ("DH1001", "08:12", "Nguyễn Văn An", "0901111222", "Zalo", 450000, 0, 450000, "00000121"),
    ("DH1002", "08:30", "Trần Thị Bình", "0902222333", "Facebook", 1250000, 50000, 1200000, "00000122"),
    ("DH1003", "09:05", "Lê Minh Châu", "0912345678", "Zalo", 320000, 0, 320000, ""),
    ("DH1004", "09:20", "Phạm Quốc Dũng", "0904444555", "Website", 890000, 0, 890000, ""),
    ("DH1005", "09:48", "Hoàng Thu Hà", "0905555666", "Zalo", 560000, 0, 560000, "00000123"),
    ("DH1006", "10:15", "Võ Thanh Em", "0906666777", "Zalo", 250000, 0, 250000, ""),
    ("DH1007", "10:40", "Đỗ Gia Huy", "0907777888", "Facebook", 1500000, 0, 1500000, "00000124"),
    ("DH1008", "11:02", "Công ty TNHH Nội thất Gỗ Việt", "0243888999", "Điện thoại", 5400000, 0, 5400000, "00000126"),
    ("DH1009", "11:25", "Bùi Ngọc Lan", "0909999000", "Zalo", 780000, 0, 780000, "00000125"),
    ("DH1010", "13:10", "Đặng Hữu Phước", "0911222333", "Website", 350000, 0, 350000, ""),
    ("DH1011", "13:45", "Ngô Thị Kim", "0913444555", "Zalo", 620000, 0, 620000, ""),
    ("DH1012", "14:05", "Công ty CP Sự kiện Ánh Dương", "0283777666", "Email", 3240000, 0, 3240000, "00000127"),
    ("DH1013", "14:30", "Trịnh Văn Long", "0915666777", "Zalo", 1150000, 0, 1150000, ""),
    ("DH1014", "14:31", "Trịnh Văn Long", "0915666777", "Zalo", 400000, 0, 400000, ""),
    ("DH1015", "15:20", "Lý Thanh Tâm", "0917888999", "Facebook", 2100000, 0, 2100000, "00000128"),
    ("DH1016", "16:00", "Mai Phương Thảo", "0918000111", "Website", 990000, 0, 990000, ""),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(ARIAL_B if bold else ARIAL), size)


def vnd(value: int, sep: str = ".") -> str:
    return f"{value:,}".replace(",", sep)


def status_bar(d: ImageDraw.ImageDraw, width: int, color="black", clock="17:42") -> None:
    d.text((40, 18), clock, font=font(26, True), fill=color)
    d.rectangle([width - 90, 24, width - 44, 44], outline=color, width=2)
    d.rectangle([width - 87, 27, width - 56, 41], fill=color)
    for i, h in enumerate((8, 12, 16, 20)):
        d.rectangle([width - 150 + i * 9, 44 - h, width - 144 + i * 9, 44], fill=color)


def footer(d: ImageDraw.ImageDraw, width: int, height: int) -> None:
    d.text((width // 2, height - 30), NOTE, font=font(18), fill=(170, 170, 170), anchor="mm")


def check_icon(d: ImageDraw.ImageDraw, cx: int, cy: int, r: int, color) -> None:
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=color)
    d.line([(cx - r * 0.45, cy), (cx - r * 0.1, cy + r * 0.38), (cx + r * 0.5, cy - r * 0.35)], fill="white", width=max(6, r // 6))


def style_sen_vang(t: dict) -> Image.Image:
    """App ngân hàng kiểu cũ: dòng 'Nhãn: giá trị'."""
    w, h = 750, 1500
    im = Image.new("RGB", (w, h), (244, 246, 245))
    d = ImageDraw.Draw(im)
    green = (0, 112, 74)
    d.rectangle([0, 0, w, 200], fill=green)
    status_bar(d, w, "white", t["clock"])
    d.text((40, 110), "Sen Vàng Bank", font=font(44, True), fill="white")
    d.rounded_rectangle([30, 240, w - 30, 1380], 24, fill="white")
    check_icon(d, w // 2, 340, 56, green)
    d.text((w // 2, 450), "Chuyển tiền thành công", font=font(36, True), fill=(30, 30, 30), anchor="mm")
    d.text((w // 2, 540), f"{vnd(t['amount'], ',')} VND", font=font(64, True), fill=green, anchor="mm")
    y = 660
    for label, value in (("Người chuyển", t["payer"]), ("Tài khoản nhận", SHOP_ACCOUNT), ("Tên người nhận", SHOP),
                         ("Nội dung", t["content"]), ("Mã giao dịch", t["txid"]), ("Thời gian", t["when"])):
        d.text((50, y), f"{label}: {value}", font=font(27), fill=(45, 45, 45))
        y += 96
    footer(d, w, h)
    return im


def style_bien_xanh(t: dict) -> Image.Image:
    """App ngân hàng mới: nhãn bên trái, giá trị căn phải."""
    w, h = 750, 1560
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    blue = (20, 80, 163)
    status_bar(d, w, "black", t["clock"])
    d.text((w // 2, 110), "Biển Xanh Bank", font=font(34, True), fill=blue, anchor="mm")
    check_icon(d, w // 2, 260, 70, (16, 160, 90))
    d.text((w // 2, 385), "Giao dịch thành công", font=font(38, True), fill=(20, 20, 20), anchor="mm")
    d.text((w // 2, 475), f"- {vnd(t['amount'], ',')} VND", font=font(60, True), fill=blue, anchor="mm")
    y = 600
    for label, value in (("Tài khoản nguồn", t.get("source", "1903 4567 8901")), ("Tên người nhận", SHOP),
                         ("Tài khoản nhận", SHOP_ACCOUNT), ("Ngân hàng nhận", "Sen Vàng Bank"),
                         ("Nội dung", t["content"]), ("Mã giao dịch", t["txid"]), ("Thời gian", t["when"])):
        d.text((50, y), label, font=font(28), fill=(110, 110, 110))
        d.text((w - 50, y), value, font=font(24 if len(value) > 18 else 28, True), fill=(20, 20, 20), anchor="ra")
        d.line([(50, y + 62), (w - 50, y + 62)], fill=(230, 230, 230), width=2)
        y += 100
    d.rounded_rectangle([50, h - 190, w - 50, h - 100], 45, fill=blue)
    d.text((w // 2, h - 145), "Về trang chủ", font=font(32, True), fill="white", anchor="mm")
    footer(d, w, h)
    return im


def style_vi_may(t: dict) -> Image.Image:
    """Ví điện tử: nhãn nhỏ nằm trên, giá trị in đậm bên dưới."""
    w, h = 750, 1500
    im = Image.new("RGB", (w, h), (252, 238, 245))
    d = ImageDraw.Draw(im)
    pink = (194, 24, 91)
    status_bar(d, w, "black", t["clock"])
    d.text((w // 2, 110), "Ví Mây", font=font(40, True), fill=pink, anchor="mm")
    d.rounded_rectangle([40, 180, w - 40, 1380], 30, fill="white")
    check_icon(d, w // 2, 290, 60, pink)
    d.text((w // 2, 400), "Chuyển tiền thành công", font=font(36, True), fill=(30, 30, 30), anchor="mm")
    d.text((w // 2, 490), f"-{vnd(t['amount'])}đ", font=font(66, True), fill=(20, 20, 20), anchor="mm")
    y = 600
    for label, value in (("Người nhận", SHOP), ("Ngân hàng", f"Sen Vàng Bank · {SHOP_ACCOUNT}"),
                         ("Lời nhắn", t["content"]), ("Mã giao dịch", t["txid"]), ("Thời gian", t["when"])):
        d.text((80, y), label, font=font(26), fill=(140, 140, 140))
        d.text((80, y + 40), value, font=font(32, True), fill=(25, 25, 25))
        y += 140
    footer(d, w, h)
    return im


def style_dong_tien(t: dict) -> Image.Image:
    """Biên lai dạng phiếu, có đường cắt răng cưa."""
    w, h = 750, 1500
    im = Image.new("RGB", (w, h), (183, 28, 28))
    d = ImageDraw.Draw(im)
    status_bar(d, w, "white", t["clock"])
    d.text((w // 2, 110), "ĐỒNG TIẾN BANK", font=font(38, True), fill="white", anchor="mm")
    d.rectangle([40, 190, w - 40, 1360], fill="white")
    for x in range(40, w - 40, 30):
        d.polygon([(x, 1360), (x + 15, 1385), (x + 30, 1360)], fill="white")
    d.text((w // 2, 270), "BIÊN LAI CHUYỂN TIỀN", font=font(36, True), fill=(183, 28, 28), anchor="mm")
    d.text((w // 2, 330), "Trạng thái: Thành công", font=font(28), fill=(46, 125, 50), anchor="mm")
    d.text((w // 2, 430), f"Số tiền: {vnd(t['amount'])} VND", font=font(46, True), fill=(20, 20, 20), anchor="mm")
    for x in range(70, w - 70, 24):
        d.line([(x, 500), (x + 12, 500)], fill=(200, 200, 200), width=2)
    y = 550
    for label, value in (("Người chuyển", t["payer"]), ("Từ tài khoản", t.get("source", "2201 8899 7766")),
                         ("Người nhận", SHOP), ("Tài khoản nhận", SHOP_ACCOUNT), ("Nội dung", t["content"]),
                         ("Số tham chiếu", t["txid"]), ("Ngày giao dịch", t["when"])):
        d.text((65, y), f"{label}: {value}", font=font(26), fill=(40, 40, 40))
        y += 92
    footer(d, w, h)
    return im


def style_internet_banking(t: dict) -> Image.Image:
    """Ảnh chụp màn hình internet banking trên máy tính."""
    w, h = 1280, 800
    im = Image.new("RGB", (w, h), (236, 239, 243))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w, 70], fill=(33, 47, 61))
    d.text((30, 20), "Biển Xanh Bank · Internet Banking", font=font(28, True), fill="white")
    d.rectangle([60, 110, w - 60, 740], fill="white", outline=(210, 214, 220))
    d.text((90, 140), "Kết quả giao dịch: Chuyển khoản thành công", font=font(30, True), fill=(16, 120, 70))
    rows = (("Tài khoản trích nợ", t.get("source", "1903 2222 1111")), ("Tên người chuyển", t["payer"]),
            ("Tài khoản ghi có", SHOP_ACCOUNT), ("Tên người hưởng", SHOP), ("Số tiền", f"{vnd(t['amount'], ',')} VND"),
            ("Nội dung", t["content"]), ("Mã giao dịch", t["txid"]), ("Thời gian", t["when"]))
    y = 210
    for label, value in rows:
        d.rectangle([90, y, 520, y + 58], fill=(246, 248, 250), outline=(220, 224, 230))
        d.rectangle([520, y, w - 90, y + 58], fill="white", outline=(220, 224, 230))
        d.text((110, y + 14), label, font=font(26), fill=(80, 80, 80))
        d.text((540, y + 14), value, font=font(26, True), fill=(20, 20, 20))
        y += 58
    d.text((w // 2, h - 28), NOTE, font=font(18), fill=(150, 150, 150), anchor="mm")
    return im


STYLES = {"sen_vang": style_sen_vang, "bien_xanh": style_bien_xanh, "vi_may": style_vi_may,
          "dong_tien": style_dong_tien, "ib": style_internet_banking}

# Ảnh khách gửi qua Zalo. Tên tệp mô phỏng tên Zalo đặt khi tải về.
TRANSFERS = [
    ("z6042817353390_8f1c2a.jpg", "sen_vang", dict(amount=450000, payer="NGUYEN VAN AN", content="DH1001 NGUYEN VAN AN", txid="FT26285104411", when="12/10/2026 08:41:07", clock="08:41")),
    ("z6042817411205_a7d3e9.jpg", "bien_xanh", dict(amount=700000, payer="TRAN THI BINH", content="DH1002 chuyen truoc", txid="FT26285219034", when="12/10/2026 09:12:45", clock="09:12")),
    ("z6042817502117_3b9e01.jpg", "vi_may", dict(amount=320000, payer="LE MINH CHAU", content="0912345678 thanh toan", txid="MM2610120933", when="09:33 - 12/10/2026", clock="09:33")),
    ("z6042817618842_c41f77.jpg", "dong_tien", dict(amount=300000, payer="HOANG THU HA", content="DH1005 lan 1", txid="DT26101200871", when="12/10/2026 10:02:19", clock="10:02")),
    ("z6042817733015_77ab02.jpg", "sen_vang", dict(amount=260000, payer="HOANG THU HA", content="DH1005 lan 2", txid="FT26285310927", when="12/10/2026 15:47:52", clock="15:47")),
    ("z6042817840331_e0d5b8.jpg", "bien_xanh", dict(amount=250000, payer="VO VAN PHUC", content="chuyen tien", txid="FT26285322108", when="12/10/2026 10:50:31", clock="10:50")),
    ("z6042817921456_19ce4d.jpg", "dong_tien", dict(amount=1500000, payer="DO GIA HUY", content="DO GIA HUY thanh toan", txid="DT26101201544", when="12/10/2026 11:05:44", clock="11:05")),
    ("z6042818034470_5d2a90.png", "ib", dict(amount=5400000, payer="CONG TY TNHH NOI THAT GO VIET", content="DH1008 TT HD 00000126", txid="FT26285400516", when="12/10/2026 11:31:02")),
    ("z6042818115902_b8e6c3.jpg", "vi_may", dict(amount=800000, payer="BUI NGOC LAN", content="DH1009 Bui Ngoc Lan", txid="MM2610121142", when="11:42 - 12/10/2026", clock="11:42")),
    ("z6042818207719_0f4ad1.jpg", "sen_vang", dict(amount=350000, payer="DANG HUU PHUOC", content="DH1010 DANG HUU PHUOC", txid="FT26285512870", when="12/10/2026 13:22:16", clock="13:22")),
    ("z6042818350288_6ac2e5.jpg", "dong_tien", dict(amount=620000, payer="NGO THI KIM", content="DH1011 Ngo Thi Kim", txid="DT26101202230", when="12/10/2026 14:01:33", clock="14:01")),
    ("z6042818462531_2e7b9f.png", "ib", dict(amount=3240000, payer="CONG TY CP SU KIEN ANH DUONG", content="DH1012 thanh toan hoa don 00000127", txid="FT26285618833", when="12/10/2026 14:20:09")),
    ("z6042818573004_d9f310.jpg", "bien_xanh", dict(amount=1550000, payer="TRINH VAN LONG", content="DH1013 DH1014", txid="FT26285702264", when="12/10/2026 14:48:55", clock="14:48")),
    ("z6042818680147_41b8ac.jpg", "vi_may", dict(amount=2100000, payer="LY THANH TAM", content="DH1015 Ly Thanh Tam", txid="MM2610121537", when="15:37 - 12/10/2026", clock="15:37")),
    ("z6042818791260_a3c6d2.jpg", "sen_vang", dict(amount=430000, payer="PHAN VAN QUANG", content="DH0987 Phan Van Quang", txid="FT26285800412", when="12/10/2026 16:12:40", clock="16:12")),
]


def make_transfers(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for name, style, data in TRANSFERS:
        im = STYLES[style](data)
        if name.startswith("z6042818207719"):
            # Ảnh chụp lại màn hình, nghiêng và bị Zalo nén mạnh.
            im = im.rotate(2.5, expand=True, fillcolor=(30, 30, 30)).filter(ImageFilter.GaussianBlur(1.1))
            im = im.resize((im.width * 3 // 4, im.height * 3 // 4))
            im.save(folder / name, quality=38)
            continue
        if name.startswith("z6042818350288"):
            # Khách cắt ảnh, mất phần số tiền phía trên.
            im = im.crop((0, 520, im.width, im.height))
        if name.endswith(".png"):
            im.save(folder / name)
        else:
            im.save(folder / name, quality=82)
    # Khách gửi lại đúng ảnh cũ.
    shutil.copy(folder / "z6042817502117_3b9e01.jpg", folder / "z6042819002388_3b9e01.jpg")
    # Khách gửi nhầm ảnh sản phẩm.
    im = Image.new("RGB", (750, 1000), (247, 235, 220))
    d = ImageDraw.Draw(im)
    for cx, cy, color in ((300, 380, (214, 69, 101)), (450, 360, (236, 120, 140)), (375, 260, (200, 50, 90))):
        d.ellipse([cx - 110, cy - 110, cx + 110, cy + 110], fill=color)
    d.rectangle([360, 470, 390, 820], fill=(70, 130, 60))
    d.text((375, 900), "Mẫu bó hoa số 3 – còn hàng không shop?", font=font(30, True), fill=(80, 50, 50), anchor="mm")
    im.save(folder / "z6042819115630_7e21aa.jpg", quality=85)


# ---------- Hóa đơn điện tử ----------

UNITS = ["", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]


def read_three(n: int, full: bool) -> str:
    hundreds, tens, ones = n // 100, n // 10 % 10, n % 10
    words = []
    if full or hundreds:
        words += [UNITS[hundreds] if hundreds else "không", "trăm"]
    if tens > 1:
        words += [UNITS[tens], "mươi"]
        if ones == 1:
            words.append("mốt")
        elif ones == 5:
            words.append("lăm")
        elif ones:
            words.append(UNITS[ones])
    elif tens == 1:
        words.append("mười")
        if ones == 5:
            words.append("lăm")
        elif ones:
            words.append(UNITS[ones])
    elif ones:
        if full or hundreds:
            words.append("lẻ")
        words.append(UNITS[ones])
    return " ".join(words)


def amount_in_words(n: int) -> str:
    groups, scales = [], ["", "nghìn", "triệu", "tỷ"]
    while n:
        groups.append(n % 1000)
        n //= 1000
    words = []
    for i in range(len(groups) - 1, -1, -1):
        if groups[i]:
            words.append(read_three(groups[i], i != len(groups) - 1))
            words.append(scales[i])
    text = " ".join(w for w in words if w).strip()
    return text[0].upper() + text[1:] + " đồng."


SELLER = dict(name="CÔNG TY TNHH HOA MỘC DEMO", tax="0109876543", address="Số 12 Phố Hoa Sữa, Phường Láng, Hà Nội",
              phone="024 3999 1234", account="0071 0001 23456 tại Sen Vàng Bank")

# số HĐ, ngày trên PDF, mã đơn, người mua, đơn vị, MST người mua, địa chỉ, dòng hàng (tên, ĐVT, SL, đơn giá trước thuế)
INVOICES = [
    ("00000121", "12/10/2026", "DH1001", "Nguyễn Văn An", "", "", "Ngõ 5 Láng Hạ, Hà Nội", [("Bó hoa hồng đỏ 20 bông", "Bó", 1, None)]),
    ("00000122", "12/10/2026", "DH1002", "Trần Thị Bình", "", "", "Cầu Giấy, Hà Nội", [("Giỏ hoa khai trương", "Giỏ", 1, None)]),
    ("00000123", "12/10/2026", "DH1005", "Hoàng Thu Hà", "", "", "Đống Đa, Hà Nội", [("Bó hoa tulip", "Bó", 1, None)]),
    ("00000124", "11/10/2026", "DH1007", "Đỗ Gia Huy", "", "", "Tây Hồ, Hà Nội", [("Kệ hoa chúc mừng", "Kệ", 1, None)]),
    ("00000126", "12/10/2026", "DH1008", "Phạm Thu Trang", "CÔNG TY TNHH NỘI THẤT GỖ VIỆT", "0108123456", "Số 88 Trần Duy Hưng, Hà Nội",
     [("Kệ hoa khai trương", "Kệ", 2, 2000000), ("Lẵng hoa để bàn", "Lẵng", 2, 500000)]),
    ("00000127", "12/10/2026", "DH1012", "Lê Hoàng Nam", "CÔNG TY CP SỰ KIỆN ÁNH DƯƠNG", "0316543201", "Số 21 Nguyễn Huệ, TP. Hồ Chí Minh",
     [("Hoa trang trí sự kiện", "Gói", 1, 2400000), ("Vận chuyển và lắp đặt", "Lần", 1, 600000)]),
    ("00000128", "12/10/2026", "DH1015", "Lý Thanh Tâm", "", "", "Hai Bà Trưng, Hà Nội", [("Bình hoa lan hồ điệp", "Bình", 1, None)]),
]
ORDER_TOTAL = {o[0]: o[7] for o in ORDERS}


def split_vat(total: int, rate: float = 0.08) -> tuple[int, int]:
    subtotal = round(total / (1 + rate))
    return subtotal, total - subtotal


def invoice_pdf(path: Path, inv: tuple) -> None:
    number, day, order_code, buyer, company, buyer_tax, address, items = inv
    total = ORDER_TOTAL[order_code]
    subtotal, vat = split_vat(total)
    if items[0][3] is None:
        items = [(items[0][0], items[0][1], 1, subtotal)]
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="tr", fontfile=str(TIMES))
    page.insert_font(fontname="tb", fontfile=str(TIMES_B))
    red, dark = (0.65, 0.1, 0.1), (0.1, 0.1, 0.1)

    def text(x, y, s, size=10, bold=False, color=dark):
        page.insert_text((x, y), s, fontname="tb" if bold else "tr", fontsize=size, color=color)

    def center(y, s, size=10, bold=False, color=dark):
        width = pymupdf.Font(fontfile=str(TIMES_B if bold else TIMES)).text_length(s, fontsize=size)
        text((595 - width) / 2, y, s, size, bold, color)

    page.draw_rect(pymupdf.Rect(20, 20, 575, 822), color=red, width=1.2)
    center(62, "HÓA ĐƠN GIÁ TRỊ GIA TĂNG", 17, True, red)
    center(80, "(VAT INVOICE)", 10)
    d, m, y = day.split("/")
    center(98, f"Ngày (Date) {d} tháng (month) {m} năm (year) {y}", 10)
    center(114, "Bản thể hiện của hóa đơn điện tử (Electronic invoice display)", 9)
    text(40, 136, f"Mã của CQT: 00F2C6A91B3D4E5F{number[-3:]}A7C", 9)
    text(420, 136, "Ký hiệu (Serial): 1C26THM", 10)
    text(420, 152, f"Số (No): {number}", 10, True, red)
    page.draw_line((30, 160), (565, 160), color=red, width=0.6)
    y = 178
    for label, value, bold in (("Đơn vị bán hàng (Seller)", SELLER["name"], True), ("Mã số thuế (Tax code)", SELLER["tax"], True),
                               ("Địa chỉ (Address)", SELLER["address"], False), ("Điện thoại (Tel)", SELLER["phone"], False),
                               ("Số tài khoản (Account No.)", SELLER["account"], False)):
        text(40, y, f"{label}: ", 10)
        text(205, y, value, 10, bold)
        y += 16
    page.draw_line((30, y - 6), (565, y - 6), color=red, width=0.6)
    y += 8
    for label, value in (("Họ tên người mua hàng (Buyer)", buyer), ("Tên đơn vị (Company name)", company),
                         ("Mã số thuế (Tax code)", buyer_tax), ("Địa chỉ (Address)", address),
                         ("Hình thức thanh toán (Payment method)", "Chuyển khoản"), ("Mã đơn hàng (Order No.)", order_code)):
        text(40, y, f"{label}: ", 10)
        text(235, y, value, 10, True)
        y += 16
    y += 6
    cols = [30, 60, 280, 330, 380, 465, 565]
    heads = ["STT", "Tên hàng hóa, dịch vụ", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"]
    page.draw_rect(pymupdf.Rect(30, y, 565, y + 22), color=red, fill=(0.98, 0.93, 0.93), width=0.6)
    for i, h in enumerate(heads):
        text(cols[i] + 4, y + 15, h, 9.5, True)
    y += 22
    rows = items + [("", "", "", "")] * (5 - len(items))
    for index, (name, unit, qty, price) in enumerate(rows, 1):
        page.draw_rect(pymupdf.Rect(30, y, 565, y + 20), color=red, width=0.4)
        if name:
            for i, value in enumerate((str(index), name, unit, str(qty), vnd(price), vnd(qty * price))):
                text(cols[i] + 4, y + 14, value, 9.5)
        y += 20
    for x in cols[1:-1]:
        page.draw_line((x, y - 20 * len(rows) - 22), (x, y), color=red, width=0.4)
    y += 18
    text(300, y, f"Cộng tiền hàng (Total amount): {vnd(subtotal)}", 10)
    y += 16
    text(40, y, "Thuế suất GTGT (VAT rate): 8%", 10)
    text(300, y, f"Tiền thuế GTGT (VAT amount): {vnd(vat)}", 10)
    y += 16
    text(300, y, f"Tổng tiền thanh toán (Total payment): {vnd(total)}", 10, True)
    y += 20
    text(40, y, f"Số tiền viết bằng chữ (Amount in words): {amount_in_words(total)}", 10)
    y += 40
    text(80, y, "Người mua hàng (Buyer)", 10, True)
    text(380, y, "Người bán hàng (Seller)", 10, True)
    page.draw_rect(pymupdf.Rect(345, y + 12, 545, y + 70), color=(0.1, 0.5, 0.2), width=1)
    text(355, y + 30, "Signature Valid", 10, True, (0.1, 0.5, 0.2))
    text(355, y + 46, f"Ký bởi: {SELLER['name']}", 8.5)
    text(355, y + 60, f"Ký ngày: {day}", 8.5)
    text(40, 800, "Tra cứu tại: hoamoc-demo.invoice.test · Mã tra cứu: HMD" + number[-3:] + "X9", 8.5)
    text(40, 812, "Hóa đơn mẫu giả lập cho khóa học AI Agent SlimAI. Không có giá trị pháp lý.", 8, False, (0.5, 0.5, 0.5))
    doc.save(path)


def make_invoices(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for inv in INVOICES:
        target = folder / f"HD_1C26THM_{inv[0]}.pdf"
        invoice_pdf(target, inv)
        if inv[0] == "00000122":
            # Hóa đơn được in ra rồi chụp lại bằng điện thoại.
            with pymupdf.open(target) as doc:
                pix = doc[0].get_pixmap(dpi=170)
            photo = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            photo = photo.rotate(-1.5, expand=True, fillcolor=(90, 80, 70)).filter(ImageFilter.GaussianBlur(0.6))
            photo.save(folder / "anh-chup-hoa-don-00000122.jpg", quality=80)
            target.unlink()
    # Kế toán tải trùng một hóa đơn hai lần.
    shutil.copy(folder / "HD_1C26THM_00000126.pdf", folder / "HD_1C26THM_00000126 (1).pdf")


# ---------- File xuất từ phần mềm bán hàng ----------

def make_export(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Đơn hàng"
    ws.append(["Mã ĐH", "Thời gian", "Khách hàng", "SĐT", "Kênh bán", "Tổng tiền hàng", "Giảm giá", "Khách cần trả",
               "Hình thức TT", "Số HĐ điện tử", "Người tạo"])
    for code, clock, name, phone, channel, goods, discount, due, inv in ORDERS:
        ws.append([code, datetime.strptime(f"12/10/2026 {clock}", "%d/%m/%Y %H:%M"), name, phone, channel, goods,
                   discount, due, "Chuyển khoản", inv or None, "thu.ngan01"])
    inv_ws = wb.create_sheet("Hóa đơn điện tử")
    inv_ws.append(["Ký hiệu", "Số hóa đơn", "Ngày hóa đơn", "Mã ĐH", "Tên người mua", "MST người mua",
                   "Tiền trước thuế", "Thuế GTGT", "Tổng tiền", "Trạng thái"])
    export_rows = {
        # Phần mềm lưu đúng; PDF lệch ở 124 (ngày) và 127 (MST). Phần mềm gõ sai tổng tiền ở 123.
        "00000121": ("12/10/2026", "DH1001", "Nguyễn Văn An", "", 450000),
        "00000122": ("12/10/2026", "DH1002", "Trần Thị Bình", "", 1200000),
        "00000123": ("12/10/2026", "DH1005", "Hoàng Thu Hà", "", 650000),
        "00000124": ("12/10/2026", "DH1007", "Đỗ Gia Huy", "", 1500000),
        "00000125": ("12/10/2026", "DH1009", "Bùi Ngọc Lan", "", 780000),
        "00000126": ("12/10/2026", "DH1008", "CÔNG TY TNHH NỘI THẤT GỖ VIỆT", "0108123456", 5400000),
        "00000127": ("12/10/2026", "DH1012", "CÔNG TY CP SỰ KIỆN ÁNH DƯƠNG", "0316543210", 3240000),
    }
    for number, (day, code, buyer, tax, total) in export_rows.items():
        subtotal, vat = split_vat(total)
        inv_ws.append(["1C26THM", number, datetime.strptime(day, "%d/%m/%Y"), code, buyer, tax or None,
                       subtotal, vat, total, "Đã cấp mã"])
    for sheet in (ws, inv_ws):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF2F5597")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(max(12, max(len(str(c.value or "")) for c in column) + 2), 34)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "dd/mm/yyyy hh:mm"
        for col in (6, 7, 8):
            ws.cell(row, col).number_format = "#,##0"
    for row in range(2, inv_ws.max_row + 1):
        inv_ws.cell(row, 3).number_format = "dd/mm/yyyy"
        for col in (7, 8, 9):
            inv_ws.cell(row, col).number_format = "#,##0"
    wb.save(path)


def main() -> None:
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "bo_du_lieu_mau_doi_soat")
    make_transfers(root / "1-anh-chuyen-khoan-zalo")
    make_invoices(root / "2-hoa-don-dien-tu")
    (root / "3-xuat-tu-phan-mem").mkdir(parents=True, exist_ok=True)
    make_export(root / "3-xuat-tu-phan-mem" / "phan-mem-ban-hang-12-10-2026.xlsx")
    print(f"Đã tạo bộ mẫu tại {root}")


if __name__ == "__main__":
    main()
