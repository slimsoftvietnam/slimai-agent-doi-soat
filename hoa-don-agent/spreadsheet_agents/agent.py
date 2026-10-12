"""Agent đối soát ảnh chuyển khoản với danh sách đơn; kèm đọc hóa đơn và xử lý bảng tính Excel."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import os
from pathlib import Path
import re
import sys
import unicodedata

ROOT = Path(__file__).resolve().parents[1]
SMART_OCR_DIR = Path(__file__).resolve().parent / "smart_ocr_base"
SUPPORTED = {".pdf", ".jpg", ".jpeg", ".png"}
RED = "FFFFC7CE"
YELLOW = "FFFFEB9C"
GREEN = "FFE2F0D9"
BLUE = "FFDDEBF7"
ORANGE = "FFF8CBAD"


def load_env() -> None:
    try:
        from dotenv import load_dotenv
    except ImportError:
        return
    load_dotenv(ROOT / ".env")


def workbook_libs():
    try:
        from openpyxl import load_workbook, Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl. Chạy .venv/bin/pip install -r spreadsheet_agents/requirements.txt") from exc
    return load_workbook, Workbook, Font, PatternFill, get_column_letter


def input_files(path: Path) -> list[Path]:
    if path.is_file():
        paths = [path]
    elif path.is_dir():
        paths = sorted(p for p in path.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED)
    else:
        raise RuntimeError(f"Không tìm thấy đầu vào: {path}")
    paths = [p for p in paths if p.suffix.lower() in SUPPORTED]
    if not paths:
        raise RuntimeError("Không tìm thấy tệp PDF/JPG/JPEG/PNG được hỗ trợ.")
    return paths


def extract_invoice(path: Path, extractor_mode: str) -> dict:
    """Run the MIT-licensed Smart-OCR pipeline, then adapt its fields."""
    load_env()
    if str(SMART_OCR_DIR) not in sys.path:
        sys.path.insert(0, str(SMART_OCR_DIR))
    if extractor_mode == "llm" and not os.getenv("GEMINI_API_KEY"):
        raise RuntimeError("--extractor llm cần GEMINI_API_KEY trong .env; dùng --extractor local nếu chưa cấu hình.")
    try:
        from main import process_file
    except ImportError as exc:
        raise RuntimeError("Thiếu thư viện OCR. Cài spreadsheet_agents/requirements.txt.") from exc
    from loguru import logger
    logger.remove()
    invoice = process_file(path, extractor_mode=extractor_mode)
    if invoice is None:
        raise RuntimeError("Không trích xuất được nội dung từ hóa đơn.")
    record = invoice.to_flat_dict()
    ocr_confidence = invoice.overall_ocr_confidence
    record["confidence"] = 1.0 if ocr_confidence is None else float(ocr_confidence) / 100
    record["field_confidence"] = dict(invoice.extraction_confidence or {})
    record["source_file"] = path.name
    record["seller_name"] = record.get("vendor_name")
    record["seller_address"] = record.get("vendor_address")
    record["seller_tax_id"] = record.get("vendor_gstin")
    record["buyer_tax_id"] = record.get("buyer_gstin")
    record["vat_amount"] = record.get("tax_amount")
    record["line_items"] = [
        {
            "description": item.description,
            "unit": (item.extra_fields or {}).get("Unit") or (item.extra_fields or {}).get("unit"),
            "quantity": item.quantity,
            "unit_price": item.unit_price,
            "line_total": item.line_total,
        }
        for item in (invoice.line_items or [])
    ]
    record["currency"] = record.get("currency") or "VND"
    record = supplement_vietnamese_receipt(path, record)
    return record


def supplement_vietnamese_receipt(path: Path, record: dict) -> dict:
    """Fill common Vietnamese digital-receipt fields from their nearby labels.

    The upstream parser is optimized for conventional tax invoices. These
    narrow, label-led rules recover receipt number/date/customer fields without
    guessing tax IDs from unrelated numeric values such as bank accounts.
    """
    try:
        if path.suffix.lower() == ".pdf":
            import pymupdf
            with pymupdf.open(path) as document:
                text = "\n".join(page.get_text("text", sort=True) for page in document)
        else:
            # Photos of printed e-invoices: OCR the whole page so the same label rules apply.
            import pytesseract
            from PIL import Image, ImageOps
            if os.getenv("TESSERACT_CMD"):
                pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
            with Image.open(path) as source:
                image = ImageOps.autocontrast(ImageOps.grayscale(ImageOps.exif_transpose(source)))
            text = pytesseract.image_to_string(image, lang=os.getenv("TESSERACT_LANG", "vie+eng"), config="--psm 4")
            if re.search(r"(?i)Tổng\s+tiền\s+thanh\s+toán", text):
                # Generic OCR money fields on photographed e-invoices are unreliable; keep only labelled values.
                record["subtotal"] = record["vat_amount"] = record["total_amount"] = None
    except Exception:
        return record
    if not text.strip():
        return record

    confidences = record.setdefault("field_confidence", {})
    # Structured e-invoices often place bilingual labels and values on the
    # same line. Recover those exact fields before applying looser heuristics.
    header_number = re.search(
        r"(?i)(?:\bSố\s*\(No\)|\bSố)\s*[:：]\s*([0-9]{3,})\b",
        text,
    )
    if header_number:
        record["invoice_number"] = header_number.group(1)
        confidences["invoice_number"] = 0.99

    vietnamese_date = re.search(
        r"(?i)Ngày(?:\s*\([^)]*\))?[\s.:-]*(\d{1,2})[\s.:-]*tháng(?:\s*\([^)]*\))?[\s.:-]*(\d{1,2})[\s.:-]*năm(?:\s*\([^)]*\))?[\s.:-]*(\d{4})",
        text,
    )
    if vietnamese_date:
        day, month, year = vietnamese_date.groups()
        record["invoice_date"] = f"{int(day):02d}/{int(month):02d}/{year}"
        confidences["invoice_date"] = 0.99

    seller = re.search(
        r"(?im)(?:Tên\s+đơn\s+vị\s+bán|Đơn\s+vị\s+bán\s+hàng)(?:\s*\([^)]*\))?\s*[:：]\s*([^\r\n]+)",
        text,
    )
    if seller and seller.group(1).strip():
        record["seller_name"] = seller.group(1).strip()
        confidences["seller_name"] = 0.99

    buyer_start = re.search(r"(?i)(?:Họ\s+tên\s+người\s+mua|Tên\s+người\s+mua)", text)
    buyer_text = text[buyer_start.start():] if buyer_start else ""
    buyer = re.search(
        r"(?im)(?:Tên\s+người\s+mua(?:\s+hàng)?|Họ\s+tên\s+người\s+mua(?:\s+hàng)?)(?:\s*\([^)]*\))?\s*[:：]\s*([^\r\n]+)",
        text,
    )
    buyer_name = buyer.group(1).strip(" .:-") if buyer else ""
    if not buyer_name:
        buyer_company = re.search(
            r"(?im)Tên\s+đơn\s+vị(?:\s*\([^)]*\))?\s*[:：]\s*([^\r\n]+)",
            buyer_text,
        )
        buyer_name = buyer_company.group(1).strip(" .:-") if buyer_company else ""
    if buyer_name:
        record["buyer_name"] = buyer_name
        confidences["buyer_name"] = 0.98

    buyer_tax = re.search(r"(?i)(?:Mã\s*số\s*thuế|MST)(?:\s*\([^)]*\))?\s*[:：]\s*([0-9]{10}(?:-?[0-9]{3})?)", buyer_text)
    if buyer_tax:
        record["buyer_tax_id"] = buyer_tax.group(1)
        confidences["buyer_tax_id"] = 0.98
    seller_section = text[:buyer_start.start()] if buyer_start else text[:2000]
    tax_id = re.search(
        r"(?i)(?:Mã\s*số\s*thuế|MST)(?:\s*\([^)]*\))?\s*[:：]\s*([0-9]{10}(?:[- ]?[0-9]{3})?)",
        seller_section,
    )
    if tax_id:
        record["seller_tax_id"] = tax_id.group(1).replace(" ", "")
        confidences["seller_tax_id"] = 0.99

    subtotal = re.search(
        r"(?im)(?:Tổng\s+tiền\s+chưa\s+có\s+thuế\s+GTGT|Cộng\s+tiền\s+h[àa]ng)\s*[^\r\n:]*[:：]?\s*([\d.,]+)",
        text,
    )
    if subtotal:
        record["subtotal"] = subtotal.group(1)
        confidences["subtotal"] = 0.98

    vat = re.search(
        r"(?im)(?:Tổng\s+số\s+tiền\s+thuế\s+giá\s+trị\s+gia\s+tăng|Tiền\s+thuế\s+GTGT|Tổng\s+tiền\s+thuế\s+GTGT)\s*[^\r\n:]*[:：]?\s*([\d.,]+)",
        text,
    )
    if vat:
        record["vat_amount"] = vat.group(1)
        confidences["vat_amount"] = 0.98

    total = re.search(
        r"(?im)(?:Tổng\s+tiền\s+thanh\s+toán(?:\s+đã\s+có\s+thuế\s+GTGT)?|Tổng\s+cộng\s+(?:tiền\s+)?thanh\s+toán)\s*[^\r\n:]*[:：]?\s*([\d.,]+)",
        text,
    )
    if total:
        record["total_amount"] = total.group(1)
        confidences["total_amount"] = 0.98

    summary_total = re.search(
        r"(?im)^\s*Tổng\s+cộng\s+tiền\s+thanh\s+toán\s*[:：]\s*([^\r\n]+)",
        text,
    )
    if summary_total:
        amounts = re.findall(r"\d[\d.,]*", summary_total.group(1))
        if len(amounts) >= 3:
            record["subtotal"], record["vat_amount"], record["total_amount"] = amounts[-3:]
            confidences.update({"subtotal": 0.98, "vat_amount": 0.98, "total_amount": 0.98})
        elif amounts:
            record["total_amount"] = amounts[-1]
            confidences["total_amount"] = 0.98

    invoice_header = re.search(
        r"(?is)\bH[ÓO]A\s+Đ[ƠO]N\s*\n\s*#?\s*([A-Z0-9][A-Z0-9/-]*)\s*\n\s*(\d{1,2}[/-]\d{1,2}[/-]\d{4})",
        text,
    )
    if invoice_header:
        if not record.get("invoice_number"):
            record["invoice_number"] = invoice_header.group(1)
            confidences["invoice_number"] = 0.98
        if not record.get("invoice_date"):
            record["invoice_date"] = invoice_header.group(2)
            confidences["invoice_date"] = 0.98

    buyer = re.search(
        r"(?is)THÔNG\s+TIN\s+KHÁCH\s+HÀNG.{0,160}?Họ\s*tên\s*\n\s*([^\r\n]+)",
        text,
    )
    if buyer and not record.get("buyer_name"):
        record["buyer_name"] = buyer.group(1).strip()
        confidences["buyer_name"] = 0.98

    if not record.get("seller_name") or re.search(r"\d{1,2}/\d{1,2}/\d{2,4}", str(record.get("seller_name"))):
        before_invoice = text.split("HÓA ĐƠN", 1)[0].splitlines()
        candidates = [
            line.strip() for line in before_invoice
            if line.strip()
            and not re.search(r"\d{1,2}/\d{1,2}/\d{2,4}|https?://", line, re.I)
            and not re.search(r"built\s+by", line, re.I)
        ]
        if candidates:
            record["seller_name"] = candidates[0]
            confidences["seller_name"] = 0.95

    payment_total = re.search(
        r"(?im)^\s*Tổng\s+thanh\s+toán\s*\n\s*([\d., ]+)\s*(?:₫|đ|VND)?",
        text,
    )
    if payment_total and not record.get("total_amount"):
        record["total_amount"] = payment_total.group(1).strip()
        confidences["total_amount"] = 0.98
    if "₫" in text or re.search(r"\bVND\b|đồng", text, re.I):
        record["currency"] = "VND"

    return record


def invoice_issues(record: dict, threshold: float) -> list[str]:
    issues = []
    confidences = record.get("field_confidence") or {}
    for label, field in (("MST người bán", "seller_tax_id"), ("Số hóa đơn", "invoice_number"), ("Ngày hóa đơn", "invoice_date"), ("Tổng tiền", "total_amount")):
        if record.get(field) in (None, ""):
            issues.append(f"Thiếu {label}")
        if float(confidences.get(field, record.get("confidence", 0)) or 0) < threshold:
            issues.append(f"Độ tin cậy {label} thấp")
    tax_id = re.sub(r"\D", "", str(record.get("seller_tax_id") or ""))
    if tax_id and len(tax_id) not in (10, 13):
        issues.append("MST người bán không có 10 hoặc 13 chữ số")
    if float(record.get("confidence", 1) or 0) < threshold:
        issues.append("Độ tin cậy OCR thấp")
    raw_date = record.get("invoice_date")
    if raw_date:
        parsed_date = normalize(raw_date, "Ngày hóa đơn")
        if not isinstance(parsed_date, (datetime, date)):
            issues.append("Ngày hóa đơn không theo YYYY-MM-DD")
    subtotal, vat, total = (normalize(record.get(key), "Tổng tiền") for key in ("subtotal", "vat_amount", "total_amount"))
    if total is not None and (subtotal is None or vat is None):
        issues.append("Thiếu tiền trước thuế hoặc thuế GTGT; chưa kiểm được phép cộng")
    if any(isinstance(value, float) and not value.is_integer() for value in (subtotal, vat, total)):
        issues.append("Số tiền có phần lẻ bất thường với VND")
    if all(isinstance(value, (int, float)) for value in (subtotal, vat, total)):
        if abs(subtotal + vat - total) > 1:
            issues.append("Tổng tiền khác tiền trước thuế + thuế GTGT")
    return list(dict.fromkeys(issues))


def remove_obvious_ocr_noise(record: dict) -> None:
    """Keep plausible partial values while dropping obvious layout/OCR debris."""
    for field in ("seller_name", "buyer_name", "seller_address", "buyer_address", "payment_method", "amount_in_words"):
        value = record.get(field)
        if not isinstance(value, str) or not value.strip():
            continue
        cleaned = value.strip()
        letters = sum(char.isalpha() for char in cleaned)
        if letters < 4 or letters / max(len(cleaned), 1) < 0.45:
            record[field] = None
            continue
        if re.search(r"(?i)(người bán hàng|thủ trưởng đơn vị|hình thức thanh toán|tên hàng hóa|mã số thuế)", cleaned):
            record[field] = None
    number = record.get("invoice_number")
    if number and not re.fullmatch(r"[A-Za-z0-9/-]{3,24}", str(number).strip()):
        record["invoice_number"] = None
    tax_id = record.get("seller_tax_id")
    if tax_id:
        digits = re.sub(r"\D", "", str(tax_id))
        if len(digits) not in (10, 13):
            record["seller_tax_id"] = None
    currency = str(record.get("currency") or "").strip().upper()
    if currency and currency not in {"VND", "VNĐ", "USD", "EUR", "JPY", "GBP"}:
        record["currency"] = None


def invoices(args) -> Path:
    load_workbook, Workbook, Font, PatternFill, get_column_letter = workbook_libs()
    files = input_files(Path(args.input).expanduser())
    rows = []
    for index, path in enumerate(files, 1):
        print(f"[{index}/{len(files)}] Đang đọc {path.name}", file=sys.stderr)
        try:
            record = extract_invoice(path, args.extractor)
            record["source_file"] = path.name
            if float(record.get("confidence", 1) or 0) < args.confidence_threshold:
                # Keep partial OCR results visible. The red review status and
                # low-confidence note signal that handwriting still needs a
                # person to confirm it; blanking all fields hid useful data.
                record["notes"] = "OCR ảnh/chữ viết tay có độ tin cậy thấp; giữ các trường đọc được để đối chiếu, không coi là dữ liệu đã xác nhận."
                remove_obvious_ocr_noise(record)
            record["issues"] = invoice_issues(record, args.confidence_threshold)
            rows.append(record)
        except Exception as exc:
            rows.append({"source_file": path.name, "notes": str(exc), "issues": ["Không trích xuất được"]})
    out = Path(args.output).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hóa đơn"
    headers = ["Tệp nguồn", "Tên người bán", "MST người bán", "Địa chỉ người bán", "Tên người mua", "MST người mua", "Địa chỉ người mua", "Số hóa đơn", "Ngày hóa đơn", "Phương thức thanh toán", "Tiền tệ", "Tiền trước thuế", "Thuế GTGT", "Tổng tiền", "Số tiền bằng chữ", "Độ tin cậy OCR", "Cần rà soát", "Ghi chú"]
    ws.append(headers)
    keys = ["source_file", "seller_name", "seller_tax_id", "seller_address", "buyer_name", "buyer_tax_id", "buyer_address", "invoice_number", "invoice_date", "payment_method", "currency", "subtotal", "vat_amount", "total_amount", "amount_in_words", "confidence", "issues", "notes"]
    for record in rows:
        values = []
        for key in keys:
            value = record.get(key)
            if key == "issues":
                value = "; ".join(value or [])
            elif key in {"invoice_date", "subtotal", "vat_amount", "total_amount"}:
                value = normalize(value, key)
            values.append(value)
        ws.append(values)
        if record.get("issues"):
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=RED)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in enumerate((28, 28, 18, 38, 30, 18, 38, 18, 16, 20, 12, 18, 16, 18, 48, 16, 48, 64), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F4E78")
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 9).number_format = "yyyy-mm-dd"
        ws.cell(row, 16).number_format = "0.0%"
        for col in (12, 13, 14):
            ws.cell(row, col).number_format = "#,##0.00"

    detail = wb.create_sheet("Chi tiết mặt hàng")
    detail_headers = ["Tệp nguồn", "Số hóa đơn", "STT", "Tên hàng hóa/dịch vụ", "Đơn vị", "Số lượng", "Đơn giá", "Thành tiền", "Trạng thái"]
    detail.append(detail_headers)
    for record in rows:
        items = record.get("line_items") or []
        if not items:
            detail.append([record.get("source_file"), record.get("invoice_number"), None, None, None, None, None, None, "OCR không tách được; cần kiểm tra ảnh"])
            continue
        for index, item in enumerate(items, 1):
            detail.append([
                record.get("source_file"), record.get("invoice_number"), index,
                item.get("description"), item.get("unit"), normalize(item.get("quantity"), "Số lượng"),
                normalize(item.get("unit_price"), "Đơn giá"), normalize(item.get("line_total"), "Thành tiền"),
                "Cần đối chiếu" if record.get("issues") else "Đã đọc",
            ])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F4E78")
    for idx, width in enumerate((28, 18, 8, 50, 14, 14, 18, 18, 36), 1):
        detail.column_dimensions[get_column_letter(idx)].width = width
    wb.save(out)
    return out


def normalize(value, column: str):
    if value is None:
        return None
    label = column.casefold()
    if isinstance(value, str):
        value = re.sub(r"\s+", " ", value).strip()
        if not value:
            return None
    if "mst" in label or "mã số thuế" in label or label in {"tax id", "tax_id"}:
        digits = re.sub(r"\D", "", str(value))
        return digits or None
    if any(term in label for term in ("ngày", "date")):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    pass
        return value
    if any(term in label for term in ("tiền", "tổng", "amount", "total", "subtotal", "vat", "thuế")):
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            cleaned = re.sub(r"[^\d,.-]", "", value)
            if cleaned:
                # Vietnamese display often uses 1.234.567,89; plain dot decimals remain supported.
                if "," in cleaned and "." in cleaned:
                    cleaned = cleaned.replace(".", "").replace(",", ".")
                elif cleaned.count(".") > 1:
                    cleaned = cleaned.replace(".", "")
                elif "." in cleaned and len(cleaned.rsplit(".", 1)[1]) == 3:
                    cleaned = cleaned.replace(".", "")
                elif "," in cleaned:
                    cleaned = cleaned.replace(",", ".")
                try:
                    return float(cleaned)
                except ValueError:
                    pass
    return value


def sheet_rows(ws):
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if not any(headers):
        raise RuntimeError(f"Trang tính {ws.title} không có hàng tiêu đề ở dòng 1.")
    return headers, list(ws.iter_rows(min_row=2, values_only=True))


def clean(args) -> Path:
    load_workbook, _, Font, PatternFill, _ = workbook_libs()
    source = Path(args.input).expanduser()
    wb = load_workbook(source)
    if args.sheet not in wb.sheetnames:
        raise RuntimeError(f"Không tìm thấy trang tính '{args.sheet}'. Hiện có: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]
    headers, rows = sheet_rows(ws)
    if "Trạng thái dữ liệu" in headers or "Lý do rà soát" in headers:
        raise RuntimeError("Trang tính đã có cột kết quả của tác tử; hãy chọn tệp nguồn hoặc trang tính khác.")
    if args.key and args.key not in headers:
        raise RuntimeError(f"Không tìm thấy cột khóa '{args.key}'.")
    key_col = headers.index(args.key) + 1 if args.key else None
    normalized_keys = []
    for row_idx, row in enumerate(rows, 2):
        values = [normalize(row[i] if i < len(row) else None, headers[i]) for i in range(len(headers))]
        for col_idx, value in enumerate(values, 1):
            ws.cell(row_idx, col_idx).value = value
            if isinstance(value, date):
                ws.cell(row_idx, col_idx).number_format = "yyyy-mm-dd"
        key = normalize(row[key_col - 1] if key_col and key_col - 1 < len(row) else None, args.key) if key_col else None
        normalized_keys.append(str(key).casefold() if key is not None else None)
    counts = Counter(k for k in normalized_keys if k)
    status_col = len(headers) + 1
    reason_col = status_col + 1
    ws.cell(1, status_col, "Trạng thái dữ liệu")
    ws.cell(1, reason_col, "Lý do rà soát")
    for offset, key in enumerate(normalized_keys, 2):
        reasons = []
        if key_col and not key:
            reasons.append("Thiếu khóa")
        if key and counts[key] > 1:
            reasons.append("Khóa bị trùng")
        missing_columns = [
            headers[col_idx]
            for col_idx in range(len(headers))
            if col_idx >= len(rows[offset - 2])
            or normalize(rows[offset - 2][col_idx], headers[col_idx]) is None
        ]
        if missing_columns:
            reasons.append("Thiếu dữ liệu: " + ", ".join(missing_columns))
        ws.cell(offset, status_col, "Cần rà soát" if reasons else "Đã chuẩn hóa")
        ws.cell(offset, reason_col, "; ".join(reasons))
        if reasons:
            for cell in ws[offset]:
                cell.fill = PatternFill("solid", fgColor=RED)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    out = Path(args.output).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def reconcile(args) -> Path:
    load_workbook, Workbook, Font, PatternFill, get_column_letter = workbook_libs()
    left_file, right_file = Path(args.left).expanduser(), Path(args.right).expanduser()
    left_wb, right_wb = load_workbook(left_file, data_only=True), load_workbook(right_file, data_only=True)
    if args.left_sheet not in left_wb.sheetnames or args.right_sheet not in right_wb.sheetnames:
        raise RuntimeError("Không tìm thấy trang tính đã chọn trong một trong hai tệp bảng tính.")
    left_ws, right_ws = left_wb[args.left_sheet], right_wb[args.right_sheet]
    lh, lr = sheet_rows(left_ws)
    rh, rr = sheet_rows(right_ws)
    if args.key not in lh or args.key not in rh:
        raise RuntimeError(f"Cột khóa '{args.key}' cần có trong cả hai trang tính.")
    columns = [c.strip() for c in args.columns.split(",") if c.strip()]
    missing = [c for c in columns if c not in lh or c not in rh]
    if missing:
        raise RuntimeError("Thiếu cột đối chiếu ở một trong hai trang tính: " + ", ".join(missing))
    li, ri = lh.index(args.key), rh.index(args.key)
    left_map, right_map = defaultdict(list), defaultdict(list)
    left_missing_key, right_missing_key = [], []
    for row in lr:
        key = normalize(row[li] if li < len(row) else None, args.key)
        if key is not None:
            left_map[str(key).casefold()].append(row)
        else:
            left_missing_key.append(row)
    for row in rr:
        key = normalize(row[ri] if ri < len(row) else None, args.key)
        if key is not None:
            right_map[str(key).casefold()].append(row)
        else:
            right_missing_key.append(row)
    wb = Workbook()
    ws = wb.active
    ws.title = "Đối chiếu"
    header = [args.key]
    for col in columns:
        header.extend([f"{col} ({left_file.stem})", f"{col} ({right_file.stem})", f"Kết quả {col}"])
    header.append("Trạng thái")
    ws.append(header)
    for key in sorted(set(left_map) | set(right_map)):
        left_rows, right_rows = left_map.get(key, []), right_map.get(key, [])
        representative = (left_rows or right_rows)[0]
        actual_key = normalize(representative[li if left_rows else ri], args.key)
        values = [actual_key]
        reasons = []
        if len(left_rows) > 1 or len(right_rows) > 1:
            reasons.append("Khóa trùng, không ghép duy nhất")
        if not left_rows:
            reasons.append(f"Chỉ có ở bảng phải ({right_file.stem})")
        if not right_rows:
            reasons.append(f"Chỉ có ở bảng trái ({left_file.stem})")
        for col in columns:
            lv = normalize(left_rows[0][lh.index(col)] if len(left_rows) == 1 else None, col)
            rv = normalize(right_rows[0][rh.index(col)] if len(right_rows) == 1 else None, col)
            if len(left_rows) != 1 or len(right_rows) != 1:
                comparison = "Không so được"
            elif lv is None and rv is None:
                comparison = "Cùng trống"
            elif lv is None or rv is None:
                comparison = "Thiếu dữ liệu"
                reasons.append(f"Thiếu dữ liệu: {col}")
            else:
                comparison = "Khớp" if lv == rv else "Khác"
            values.extend([lv, rv, comparison])
            if len(left_rows) == len(right_rows) == 1 and lv is not None and rv is not None and lv != rv:
                reasons.append(f"Sai khác: {col}")
        status = "; ".join(dict.fromkeys(reasons)) or "Khớp"
        values.append(status)
        ws.append(values)
        fill_color = RED if any(r.startswith(("Sai khác", "Khóa trùng", "Thiếu dữ liệu")) for r in reasons) else (YELLOW if reasons else GREEN)
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=fill_color)

    # Preserve records without a usable key. They cannot be paired, but should
    # remain visible and explicitly marked instead of silently disappearing.
    for side, missing_rows, source_headers in (
        ("trái", left_missing_key, lh),
        ("phải", right_missing_key, rh),
    ):
        for row in missing_rows:
            values = [None]
            for col in columns:
                source_value = normalize(row[source_headers.index(col)], col)
                pair = [source_value, None] if side == "trái" else [None, source_value]
                values.extend([pair[0], pair[1], "Thiếu khóa"])
            values.append(f"Thiếu khóa ở bảng {side}")
            ws.append(values)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=RED)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF1F4E78")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for idx, column in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(16, max(len(str(c.value or "")) for c in column) + 2), 38)
    out = Path(args.output).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def normalize_match(value) -> str:
    """Normalize Vietnamese text for cautious, explainable record matching."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFD", str(value).casefold().replace("đ", "d"))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def payment_proof_files(path: Path) -> list[Path]:
    supported = {".jpg", ".jpeg", ".png"}
    if path.is_file():
        files = [path]
    elif path.is_dir():
        files = sorted(p for p in path.iterdir() if p.is_file() and p.suffix.lower() in supported)
    else:
        raise RuntimeError(f"Không tìm thấy ảnh chuyển khoản: {path}")
    files = [p for p in files if p.suffix.lower() in supported]
    if not files:
        raise RuntimeError("Không tìm thấy ảnh JPG/JPEG/PNG trong đầu vào ảnh chuyển khoản.")
    return files


def labelled_value(text: str, labels: tuple[str, ...]) -> str | None:
    pattern = "|".join(re.escape(label) for label in labels)
    match = re.search(rf"(?im)^\s*(?:{pattern})\s*[:：-]?\s*([^\r\n]+)", text)
    return match.group(1).strip() if match else None


def payment_amount(text: str) -> float | None:
    """Select the most likely VND amount without mistaking account, phone or reference numbers."""
    candidates = []
    # Thousands groups use "." or ","; spaces are not accepted because account numbers are written "1903 4567 8901".
    for match in re.finditer(r"(?<![A-Za-z0-9.,])(\d{1,3}(?:[.,]\d{3})+|\d{4,})(?![0-9])", text):
        token = match.group(1)
        digits = re.sub(r"\D", "", token)
        if digits.startswith("0"):
            continue  # account numbers and phone numbers, never an amount
        value = int(digits)
        if value < 1000 or value > 10**10:
            continue
        before = text[max(0, match.start() - 45): match.start()].casefold()
        after = text[match.end(): match.end() + 8].casefold()
        score = 0
        if re.match(r"\s*(vnd|vnđ|₫|đ\b|đồng)", after):
            score += 10
        if re.search(r"số tiền|amount|thành công|đã chuyển|tổng tiền|thanh toán", before):
            score += 6
        if re.search(r"[.,]", token):
            score += 3
        else:
            score -= 3  # long bare numbers are usually codes
        line_start = text.rfind("\n", 0, match.start()) + 1
        line_label = text[line_start: match.start()].casefold()
        if re.search(r"tài khoản|account|mã|số tham chiếu|reference|ngày|thời gian|nội dung|lời nhắn|sđt|điện thoại", line_label):
            score -= 12
        candidates.append((score, value))
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    return float(candidates[0][1]) if candidates[0][0] > 0 else None


def extract_payment_proof(path: Path) -> dict:
    """OCR a downloaded bank-transfer screenshot; never treats OCR as final proof."""
    load_env()
    try:
        import pytesseract
        from PIL import Image, ImageEnhance, ImageOps
    except ImportError as exc:
        raise RuntimeError("Thiếu pytesseract/Pillow. Cài spreadsheet_agents/requirements.txt.") from exc
    if os.getenv("TESSERACT_CMD"):
        pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
    language = os.getenv("TESSERACT_LANG", "vie+eng")
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        gray = ImageOps.autocontrast(ImageOps.grayscale(image))
        if max(gray.size) < 1800:
            gray = gray.resize((gray.width * 2, gray.height * 2))
        enhanced = ImageEnhance.Contrast(gray).enhance(1.45)
        texts = []
        for variant, psm in ((gray, 6), (enhanced, 11)):
            try:
                texts.append(pytesseract.image_to_string(variant, lang=language, config=f"--psm {psm}"))
            except pytesseract.TesseractError:
                texts.append(pytesseract.image_to_string(variant, lang="eng", config=f"--psm {psm}"))
    text = max(texts, key=lambda value: len(re.sub(r"\s", "", value)), default="").strip()
    transaction_id = labelled_value(text, ("Mã giao dịch", "Mã GD", "Transaction ID", "Reference", "Số tham chiếu"))
    payer = labelled_value(text, ("Người chuyển", "Tên người gửi", "Chủ tài khoản", "Từ tài khoản", "From"))
    content = labelled_value(text, ("Nội dung", "Nội dung chuyển khoản", "Lời nhắn", "Description", "Message"))
    transferred_at = labelled_value(text, ("Thời gian", "Ngày giao dịch", "Transaction time", "Ngày"))
    if transferred_at and not re.search(r"\d{1,2}[/:]\d{2}", transferred_at):
        transferred_at = None
    return {
        "source_file": path.name,
        "source_path": str(path),
        "file_hash": hashlib.sha256(path.read_bytes()).hexdigest(),
        "raw_text": text,
        "amount": payment_amount(text),
        "transaction_id": transaction_id,
        "payer_name": payer,
        "content": content,
        "transferred_at": transferred_at,
    }


PAYMENT_COLUMNS = {
    "order_code": ("ma don", "ma don hang", "ma dh", "so don", "ma", "order", "order id", "order code"),
    "customer": ("khach hang", "ten khach", "ten khach hang", "ho ten", "ho va ten", "ten", "customer"),
    "amount_due": ("phai thu", "can thu", "khach can tra", "khach phai tra", "so tien", "so tien phai thu", "thanh tien", "tong tien", "tong", "gia tri don", "amount", "total"),
    "phone": ("so dien thoai", "dien thoai", "sdt", "so dt", "phone"),
}
LARGE_ORDER = 2_000_000  # from this amount, check the bank statement before delivery
STATUS_ORDER = ("Đã trả đủ", "Trả thiếu", "Chưa trả", "Cần kiểm tra")
STATUS_FILL = {"Đã trả đủ": GREEN, "Trả thiếu": ORANGE, "Chưa trả": RED, "Cần kiểm tra": YELLOW}


def parse_money(value) -> float | None:
    """Read VND amounts such as 1500000, '1.500.000đ' or '1,500,000'."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.]", "", str(value))
    if not text:
        return None
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", text):
        return float(re.sub(r"\D", "", text))
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def money(value) -> str:
    return f"{value:,.0f}".replace(",", ".") + "đ"


def payment_columns(headers: list[str], args) -> dict[str, str]:
    """Use the columns the user named; otherwise recognise common Vietnamese headers."""
    explicit = {
        "order_code": args.order_code_column,
        "customer": args.customer_column,
        "amount_due": args.amount_column,
        "phone": args.phone_column,
    }
    normalized = {normalize_match(header): header for header in headers if header}
    found = {}
    for field, aliases in PAYMENT_COLUMNS.items():
        if explicit[field]:
            if explicit[field] not in headers:
                raise RuntimeError(f"Không có cột '{explicit[field]}'. Các cột hiện có: {', '.join(h for h in headers if h)}")
            found[field] = explicit[field]
            continue
        match = next((normalized[alias] for alias in aliases if alias in normalized and normalized[alias] not in found.values()), None)
        if match:
            found[field] = match
    missing = [label for field, label in (("order_code", "mã đơn (--cot-ma-don)"), ("customer", "tên khách (--cot-khach-hang)"), ("amount_due", "số phải thu (--cot-so-tien)")) if field not in found]
    if missing:
        raise RuntimeError("Không nhận ra cột " + ", ".join(missing) + ". Các cột hiện có: " + ", ".join(h for h in headers if h))
    return found


def contains_code(code: str, raw: str) -> bool:
    compact = code.replace(" ", "")
    if len(compact) < 3:
        return False
    if compact.isdigit():
        # Pure-number codes must stand alone so they are not found inside account numbers.
        return re.search(rf"(?<!\d){re.escape(compact)}(?!\d)", raw) is not None
    return compact in raw.replace(" ", "")


def payment_match(proof: dict, orders: list[dict]) -> tuple[dict | None, int, str, list[dict]]:
    """Return (sure match, score, evidence, weak candidates). Amount alone never makes a sure match."""
    raw = normalize_match(" ".join(str(proof.get(key) or "") for key in ("raw_text", "content", "payer_name")))
    raw_digits = re.sub(r"\D", "", str(proof.get("raw_text") or ""))
    paid = proof.get("amount")
    scored = []
    for order in orders:
        score = 0
        evidence = []
        code = normalize_match(order.get("order_code"))
        phone = re.sub(r"\D", "", str(order.get("phone") or ""))
        name = normalize_match(order.get("customer"))
        due = order.get("amount_due")
        if code and contains_code(code, raw):
            score += 100
            evidence.append("có mã đơn")
        if phone and len(phone) >= 9 and phone[-9:] in raw_digits:
            score += 85
            evidence.append("có số điện thoại")
        if name and len(name) >= 4 and name in raw:
            score += 60
            evidence.append("có tên khách")
        if due is not None and paid is not None and abs(float(due) - float(paid)) <= 1:
            score += 35
            evidence.append("đúng số tiền")
        scored.append((score, order, ", ".join(evidence)))
    scored.sort(key=lambda item: item[0], reverse=True)
    if not scored or scored[0][0] == 0:
        return None, 0, "Không thấy mã đơn, SĐT, tên khách hay số tiền nào khớp", []
    top = scored[0][0]
    tied = [order for score, order, _ in scored if score == top]
    if top >= 70 and (len(scored) == 1 or top - scored[1][0] >= 20):
        return scored[0][1], top, scored[0][2], []
    if top >= 70:
        close = [order for score, order, _ in scored if top - score < 20][:3]
        return None, top, "Nhiều đơn khớp ngang nhau", close
    # Weak evidence (name only, amount only) only points a human to at most three orders.
    candidates = tied if len(tied) <= 3 else []
    reason = f"Chỉ {scored[0][2]}, chưa đủ để ghép" if candidates else f"Chỉ {scored[0][2]}, khớp với {len(tied)} đơn"
    return None, top, reason, candidates


def shared_transfer(proof: dict, candidates: list[dict]) -> list[dict]:
    """One transfer may pay several orders: every order code must appear and the dues must add up exactly."""
    raw = normalize_match(" ".join(str(proof.get(key) or "") for key in ("raw_text", "content")))
    coded = [o for o in candidates if normalize_match(o.get("order_code")) and contains_code(normalize_match(o.get("order_code")), raw)]
    dues = [o.get("amount_due") for o in coded]
    if len(coded) < 2 or len(coded) != len(candidates) or None in dues or proof.get("amount") is None:
        return []
    return coded if abs(sum(dues) - proof["amount"]) <= 1 else []


def reminder_message(order: dict, status: str, received: float, missing: float) -> str:
    name = order.get("customer") or "anh/chị"
    code = order.get("order_code")
    if status == "Trả thiếu":
        return (f"Chào {name}, shop đã nhận {money(received)} cho đơn {code}. "
                f"Đơn còn thiếu {money(missing)}, anh/chị kiểm tra giúp shop nhé. Cảm ơn anh/chị!")
    if status == "Chưa trả":
        return (f"Chào {name}, đơn {code} ({money(order['amount_due'])}) shop chưa thấy thanh toán. "
                "Nếu anh/chị đã chuyển, gửi giúp shop ảnh chuyển khoản nhé. Cảm ơn anh/chị!")
    return ""


def order_status(order: dict) -> tuple[str, str]:
    due = order.get("amount_due")
    received = sum(proof["amount"] for proof in order["payments"])
    hints = ", ".join(dict.fromkeys(order["hints"]))
    if due is None:
        return "Cần kiểm tra", "Danh sách đơn thiếu số phải thu; bổ sung rồi chạy lại"
    if hints and received < due - 1:
        return "Cần kiểm tra", f"Mở ảnh {hints} để xác nhận"
    if not order["payments"]:
        return "Chưa trả", "Nhắc khách thanh toán"
    if received > due + 1:
        return "Cần kiểm tra", f"Nhận dư {money(received - due)}; xác minh hoàn tiền hoặc trừ đơn sau"
    if received < due - 1:
        return "Trả thiếu", f"Nhắc khách chuyển thêm {money(due - received)}"
    if due >= LARGE_ORDER:
        return "Đã trả đủ", "Đủ theo ảnh; đơn lớn, xem sao kê trước khi giao"
    return "Đã trả đủ", "Xác nhận đơn"


def unique_output(path: Path) -> Path:
    if not path.exists():
        return path
    for n in range(2, 1000):
        candidate = path.with_name(f"{path.stem}-{n}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Không tạo được tên tệp mới cho {path}")


def reconcile_payments(args) -> Path:
    load_workbook, Workbook, Font, PatternFill, get_column_letter = workbook_libs()
    order_file = Path(args.orders).expanduser()
    wb_orders = load_workbook(order_file, data_only=True)
    if args.sheet and args.sheet not in wb_orders.sheetnames:
        raise RuntimeError(f"Không tìm thấy trang tính '{args.sheet}'. Hiện có: {', '.join(wb_orders.sheetnames)}")
    source_sheet = wb_orders[args.sheet] if args.sheet else wb_orders[wb_orders.sheetnames[0]]
    headers, rows = sheet_rows(source_sheet)
    columns = payment_columns(headers, args)
    indexes = {field: headers.index(header) for field, header in columns.items()}
    orders = []
    for row_number, row in enumerate(rows, 2):
        values = {field: row[index] if index < len(row) else None for field, index in indexes.items()}
        if all(value in (None, "") for value in values.values()):
            continue
        orders.append({
            "row_number": row_number,
            "order_code": values.get("order_code"),
            "customer": values.get("customer"),
            "phone": values.get("phone"),
            "amount_due": parse_money(values.get("amount_due")),
            "payments": [],
            "hints": [],
        })
    if not orders:
        raise RuntimeError(f"Trang '{source_sheet.title}' không có đơn nào.")

    proofs = []
    seen_hashes, seen_transactions = {}, {}
    files = payment_proof_files(Path(args.images).expanduser())
    for index, path in enumerate(files, 1):
        print(f"[{index}/{len(files)}] Đang đọc ảnh {path.name}", file=sys.stderr)
        try:
            proof = extract_payment_proof(path)
        except Exception as exc:
            proofs.append({"source_file": path.name, "result": f"Không đọc được ảnh: {exc}", "fill": RED})
            continue
        transaction_key = normalize_match(proof.get("transaction_id"))
        duplicate_of = seen_hashes.get(proof["file_hash"]) or (seen_transactions.get(transaction_key) if transaction_key else None)
        seen_hashes.setdefault(proof["file_hash"], path.name)
        if transaction_key:
            seen_transactions.setdefault(transaction_key, path.name)
        order, score, evidence, candidates = payment_match(proof, orders)
        proof.update({"match_score": score, "match_evidence": evidence})
        if duplicate_of:
            proof.update(result=f"Trùng với ảnh {duplicate_of}; không cộng tiền", fill=RED)
        elif proof.get("amount") is None:
            target = [order] if order else candidates
            for item in target:
                item["hints"].append(f"{path.name} (có {evidence.split(',')[0].removeprefix('có ')} nhưng không đọc được số tiền)" if order
                                     else f"{path.name} ({evidence.lower()})")
            proof.update(matched_order=", ".join(str(o.get("order_code")) for o in target) or None,
                         result="Không đọc được số tiền; mở ảnh kiểm tra" if proof.get("transaction_id") or target
                         else "Không đọc được số tiền; có thể không phải biên nhận", fill=RED)
        elif order:
            order["payments"].append(proof)
            proof.update(matched_order=order.get("order_code"), result="Đã ghép", fill=GREEN)
        elif shared := shared_transfer(proof, candidates):
            evidence = f"1 lần chuyển cho {len(shared)} đơn: có mã các đơn, tổng tiền khớp"
            for item in shared:
                item["payments"].append({"source_file": path.name, "amount": item["amount_due"], "match_evidence": evidence})
            proof.update(matched_order=", ".join(str(o.get("order_code")) for o in shared), match_evidence=evidence,
                         result=f"Đã ghép, chia cho {len(shared)} đơn", fill=GREEN)
        else:
            for item in candidates:
                item["hints"].append(f"{path.name} ({evidence.lower()})")
            proof.update(matched_order=", ".join(str(o.get("order_code")) for o in candidates) or None,
                         result=("Gợi ý, cần người xác nhận" if candidates else "Chưa ghép được; xem ảnh và tìm đơn"), fill=YELLOW)
        proofs.append(proof)

    out = unique_output(Path(args.output).expanduser())
    if out.resolve() == order_file.resolve():
        raise RuntimeError("Tệp kết quả trùng với danh sách đơn gốc; chọn tên khác.")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Tóm tắt"
    ws = wb.create_sheet("Đơn hàng")
    ws.append(["Mã đơn", "Khách hàng", "Số điện thoại", "Phải thu", "Đã nhận theo ảnh", "Còn thiếu", "Trạng thái",
               "Ảnh đã ghép", "Căn cứ ghép", "Việc cần làm", "Tin nhắn nhắc (soạn sẵn, chưa gửi)"])
    groups = defaultdict(list)
    for order in orders:
        status, action = order_status(order)
        received = sum(proof["amount"] for proof in order["payments"])
        due = order.get("amount_due")
        missing = max(due - received, 0) if due is not None else None
        groups[status].append((order, received, missing))
        ws.append([
            order.get("order_code"), order.get("customer"), order.get("phone"), due, received, missing, status,
            ", ".join(proof["source_file"] for proof in order["payments"]),
            "; ".join(dict.fromkeys(proof["match_evidence"] for proof in order["payments"])),
            action, reminder_message(order, status, received, missing or 0),
        ])
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=STATUS_FILL[status])

    detail = wb.create_sheet("Ảnh chuyển khoản")
    detail.append(["Tệp ảnh", "Số tiền đọc được", "Nội dung", "Người chuyển", "Mã giao dịch", "Thời gian",
                   "Ghép vào đơn", "Căn cứ", "Kết quả", "Chữ đọc được trên ảnh"])
    for proof in proofs:
        detail.append([
            proof.get("source_file"), proof.get("amount"), proof.get("content"), proof.get("payer_name"),
            proof.get("transaction_id"), proof.get("transferred_at"), proof.get("matched_order"),
            proof.get("match_evidence"), proof.get("result"), proof.get("raw_text"),
        ])
        for cell in detail[detail.max_row]:
            cell.fill = PatternFill("solid", fgColor=proof["fill"])

    total_due = sum(order["amount_due"] or 0 for order in orders)
    total_received = sum(proof["amount"] for order in orders for proof in order["payments"])
    summary.append([f"Đối soát chuyển khoản · {date.today():%d/%m/%Y}"])
    summary.append([f"{len(orders)} đơn · {len(proofs)} ảnh · phải thu {money(total_due)} · đã nhận theo ảnh {money(total_received)}"])
    summary.append([])
    summary.append(["Trạng thái", "Số đơn", "Đơn"])
    lines = []
    for status in STATUS_ORDER:
        items = groups.get(status, [])
        if status == "Trả thiếu":
            names = [f"{o['order_code']} {o['customer']} (thiếu {money(m)})" for o, _, m in items]
        elif status == "Chưa trả":
            names = [f"{o['order_code']} {o['customer']} ({money(o['amount_due'])})" for o, _, _ in items]
        else:
            names = [f"{o['order_code']} {o['customer']}" for o, _, _ in items]
        summary.append([status, len(items), "; ".join(names)])
        for cell in summary[summary.max_row]:
            cell.fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        lines.append(f"  {status}: {len(items)}" + (f" — {'; '.join(names)}" if names else ""))
    by_result = Counter("Trùng" if p["fill"] == RED and "Trùng" in p["result"] else
                        "Đã ghép" if p["fill"] == GREEN else "Cần xem" for p in proofs)
    summary.append([])
    summary.append(["Ảnh", len(proofs), f"Đã ghép {by_result['Đã ghép']} · Trùng {by_result['Trùng']} · Cần xem {by_result['Cần xem']}"])
    summary.append([])
    summary.append(["Lưu ý: số tiền lấy từ ảnh khách gửi. Đơn giá trị lớn cần xem sao kê ngân hàng trước khi giao. "
                    "Agent không gửi tin nhắn; cột tin nhắn chỉ là bản soạn sẵn."])
    summary["A1"].font = Font(bold=True, size=14)
    for cell in summary[4]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 9
    summary.column_dimensions["C"].width = 110

    for sheet in (ws, detail):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor="FF1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for idx, column in enumerate(sheet.columns, 1):
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(14, max(len(str(cell.value or "")) for cell in column) + 2), 52)
    for row in range(2, ws.max_row + 1):
        for col in (4, 5, 6):
            ws.cell(row, col).number_format = "#,##0"
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 2).number_format = "#,##0"
    wb.save(out)
    print("Kết quả đối soát:\n" + "\n".join(lines), file=sys.stderr)
    return out


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Tác tử đọc hóa đơn, đối soát thanh toán và xử lý bảng tính.")
    sub = p.add_subparsers(dest="command", required=True)
    inv = sub.add_parser("hoa-don", help="Đọc hóa đơn PDF/ảnh và xuất bảng tính")
    inv.add_argument("input", metavar="DUONG_DAN", help="Tệp hóa đơn hoặc thư mục chứa hóa đơn")
    inv.add_argument("--dau-ra", dest="output", required=True, help="Đường dẫn tệp Excel đầu ra")
    inv.add_argument("--nguong-do-tin-cay", dest="confidence_threshold", type=float, default=0.75, help="Ngưỡng OCR từ 0 đến 1; mặc định 0.75")
    def parse_extractor(value: str) -> str:
        aliases = {"cuc-bo": "local", "cục-bộ": "local", "gemini": "llm"}
        if value not in aliases:
            raise argparse.ArgumentTypeError("Chọn 'cuc-bo' hoặc 'gemini'.")
        return aliases[value]

    # argparse also applies `type` to string defaults, so keep this in the
    # public CLI vocabulary and let parse_extractor map it to the upstream mode.
    inv.add_argument("--che-do", dest="extractor", type=parse_extractor, default="cuc-bo", metavar="{cuc-bo,gemini}", help="Mặc định xử lý tại máy; Gemini cần GEMINI_API_KEY")
    inv.set_defaults(func=invoices)
    cleaner = sub.add_parser("lam-sach", help="Chuẩn hóa dữ liệu và đánh dấu hàng cần rà soát")
    cleaner.add_argument("input", metavar="TAP_TIN")
    cleaner.add_argument("--dau-ra", dest="output", required=True, help="Đường dẫn tệp Excel đầu ra")
    cleaner.add_argument("--trang-tinh", dest="sheet", default="Sheet1", help="Tên trang tính cần xử lý")
    cleaner.add_argument("--khoa", dest="key", help="Tên cột khóa tùy chọn để tìm dòng trùng")
    cleaner.set_defaults(func=clean)
    rec = sub.add_parser("doi-chieu", help="Đối chiếu hai bảng theo cột khóa")
    rec.add_argument("left", metavar="BANG_TRAI")
    rec.add_argument("right", metavar="BANG_PHAI")
    rec.add_argument("--trang-trai", dest="left_sheet", required=True, help="Tên trang tính trong bảng trái")
    rec.add_argument("--trang-phai", dest="right_sheet", required=True, help="Tên trang tính trong bảng phải")
    rec.add_argument("--khoa", dest="key", required=True, help="Tên cột dùng để ghép dòng")
    rec.add_argument("--cot", dest="columns", required=True, help="Các cột cần so sánh, cách nhau bằng dấu phẩy")
    rec.add_argument("--dau-ra", dest="output", required=True, help="Đường dẫn tệp Excel kết quả")
    rec.set_defaults(func=reconcile)
    pay = sub.add_parser("doi-soat-thanh-toan", help="Dò ảnh chuyển khoản với danh sách đơn: ai trả đủ, trả thiếu, chưa trả")
    pay.add_argument("images", metavar="ANH_CHUYEN_KHOAN", help="Ảnh hoặc thư mục ảnh JPG/JPEG/PNG đã tải từ Zalo")
    pay.add_argument("orders", metavar="DANH_SACH_DON", help="Tệp Excel danh sách đơn hàng")
    pay.add_argument("--trang-tinh", dest="sheet", help="Tên trang tính danh sách đơn; mặc định trang đầu tiên")
    pay.add_argument("--cot-ma-don", dest="order_code_column", help="Cột mã đơn; mặc định tự nhận (Mã đơn, Mã ĐH...)")
    pay.add_argument("--cot-khach-hang", dest="customer_column", help="Cột tên khách; mặc định tự nhận (Khách hàng, Họ tên...)")
    pay.add_argument("--cot-so-tien", dest="amount_column", help="Cột số phải thu; mặc định tự nhận (Phải thu, Số tiền, Tổng tiền...)")
    pay.add_argument("--cot-dien-thoai", dest="phone_column", help="Cột số điện thoại; mặc định tự nhận nếu có (SĐT, Số điện thoại...)")
    pay.add_argument("--dau-ra", dest="output", default=f"output/doi-soat-{date.today():%Y-%m-%d}.xlsx",
                     help="Tệp Excel kết quả; trùng tên thì tự thêm số, không ghi đè")
    pay.set_defaults(func=reconcile_payments)
    return p


def main() -> int:
    args = parser().parse_args()
    try:
        if args.command == "hoa-don" and not 0 <= args.confidence_threshold <= 1:
            raise RuntimeError("--nguong-do-tin-cay phải nằm trong khoảng 0 đến 1.")
        output = args.func(args)
        print(f"Đã lưu bảng tính: {output}")
        return 0
    except Exception as exc:
        print(f"Lỗi: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
